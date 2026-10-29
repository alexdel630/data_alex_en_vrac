# -*- coding: utf-8 -*-
"""Génère modele_presentation.pptx et kit_schemas.pptx (formes natives, recolorables).
Modifier la charte ci-dessous puis relancer : python build.py
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ===================== CHARTE =====================
PRINCIPAL = "1F2A44"   # bleu nuit
ACCENT = "2F6FED"      # bleu
GRIS = "667085"
FOND = "F5F7FA"
BLANC = "FFFFFF"
POLICE = "Calibri"
# Teintes secondaires dérivées (schémas multi-niveaux)
NUANCES = [PRINCIPAL, ACCENT, "5B8DEF", "9DB8F5", "C9D7FA"]
# ==================================================

ICI = Path(__file__).resolve().parent
W, H = Inches(13.333), Inches(7.5)


def rgb(h):
    return RGBColor.from_string(h)


def forme(sl, kind, x, y, w, h, fill=None, line=None, texte=None, taille=14,
          couleur=BLANC, gras=False, align=PP_ALIGN.CENTER, ancre=MSO_ANCHOR.MIDDLE):
    s = sl.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if texte is not None:
        remplir(s.text_frame, texte, taille, couleur, gras, align, ancre)
    return s


def remplir(tf, texte, taille, couleur, gras=False, align=PP_ALIGN.LEFT, ancre=MSO_ANCHOR.TOP):
    tf.word_wrap = True
    tf.vertical_anchor = ancre
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Inches(0.1))
    lignes = texte if isinstance(texte, list) else [texte]
    for i, l in enumerate(lignes):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = l
        f = r.font; f.size = Pt(taille); f.name = POLICE; f.bold = gras; f.color.rgb = rgb(couleur)


def texte(sl, x, y, w, h, t, taille=16, couleur=PRINCIPAL, gras=False, align=PP_ALIGN.LEFT,
          ancre=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    remplir(tb.text_frame, t, taille, couleur, gras, align, ancre)
    return tb


def trait(sl, x1, y1, x2, y2, couleur=GRIS, ep=1.5, fleche=False):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(couleur); c.line.width = Pt(ep)
    if fleche:
        ln = c.line._get_or_add_ln()
        from pptx.oxml.ns import qn
        from lxml import etree
        e = etree.SubElement(ln, qn("a:tailEnd")); e.set("type", "triangle")
    return c


def nouvelle(prs, fond=FOND):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.background.fill.solid(); sl.background.fill.fore_color.rgb = rgb(fond)
    return sl


def entete(sl, titre, sous=None, num=None):
    forme(sl, MSO_SHAPE.RECTANGLE, 0.6, 0.55, 0.08, 0.7, fill=ACCENT)
    texte(sl, 0.8, 0.45, 11, 0.6, titre, 28, PRINCIPAL, True)
    if sous:
        texte(sl, 0.8, 1.0, 11, 0.4, sous, 14, GRIS)
    if num is not None:
        texte(sl, 11.9, 6.95, 1, 0.35, str(num), 10, GRIS, align=PP_ALIGN.RIGHT)
        texte(sl, 0.6, 6.95, 6, 0.35, "Nom de l'organisation  |  Titre de la présentation", 10, GRIS)


def nouvelle_prs():
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    return prs


# ===================== MODELE PRESENTATION =====================
def modele():
    prs = nouvelle_prs(); n = 0

    # 1 Titre
    sl = nouvelle(prs, PRINCIPAL); n += 1
    forme(sl, MSO_SHAPE.RECTANGLE, 0, 0, 0.35, 7.5, fill=ACCENT)
    forme(sl, MSO_SHAPE.OVAL, 8.6, 1.2, 4.6, 4.6, fill="2A3758")
    forme(sl, MSO_SHAPE.OVAL, 10.9, 4.4, 2.2, 2.2, fill=ACCENT)
    texte(sl, 1.0, 2.2, 8.5, 1.6, "Titre de la présentation sur deux lignes maximum", 40, BLANC, True)
    texte(sl, 1.0, 3.9, 8.5, 0.6, "Sous-titre : contexte, objectif ou public visé", 20, "C9D7FA")
    forme(sl, MSO_SHAPE.RECTANGLE, 1.0, 4.75, 1.2, 0.06, fill=ACCENT)
    texte(sl, 1.0, 5.0, 8, 0.8, ["Prénom Nom  |  Fonction", "Lieu, JJ mois AAAA"], 14, "C9D7FA")

    # 2 Sommaire
    sl = nouvelle(prs); n += 1; entete(sl, "Sommaire", num=n)
    parties = ["Contexte et enjeux", "Diagnostic", "Propositions", "Plan d'action", "Prochaines étapes"]
    for i, p in enumerate(parties):
        y = 1.7 + i * 0.95
        forme(sl, MSO_SHAPE.OVAL, 1.2, y, 0.7, 0.7, fill=ACCENT if i == 0 else PRINCIPAL,
              texte=f"{i+1:02d}", taille=16, gras=True)
        texte(sl, 2.2, y + 0.08, 8, 0.55, p, 22, PRINCIPAL, i == 0)
        trait(sl, 2.2, y + 0.8, 11.5, y + 0.8, "D0D5DD", 0.75)

    # 3 Séparateur
    sl = nouvelle(prs, ACCENT); n += 1
    texte(sl, 1.0, 1.6, 4, 2.2, "01", 120, BLANC, True)
    texte(sl, 1.0, 4.0, 10, 0.9, "Titre de la partie", 40, BLANC, True)
    texte(sl, 1.0, 4.9, 10, 0.6, "Une phrase qui annonce ce que contient cette partie.", 18, "E3EBFD")
    forme(sl, MSO_SHAPE.RECTANGLE, 1.0, 5.7, 1.2, 0.06, fill=BLANC)

    # 4 Texte + visuel
    sl = nouvelle(prs); n += 1; entete(sl, "Texte et visuel", "Un message principal, illustré", n)
    texte(sl, 0.8, 1.8, 5.8, 0.6, "Message clé en une phrase", 22, ACCENT, True)
    texte(sl, 0.8, 2.5, 5.8, 3.8, ["• Premier argument, formulé simplement.",
          "• Deuxième argument, avec un fait ou un chiffre.", "• Troisième argument, orienté action.",
          "", "Paragraphe de conclusion de deux lignes au maximum pour garder la diapo lisible."], 16, PRINCIPAL)
    v = forme(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 7.2, 1.7, 5.3, 4.8, fill="E3EBFD",
              texte="Zone visuel\n(insérer image, picto ou graphique)", taille=16, couleur=GRIS)
    v.adjustments[0] = 0.05

    # 5 Trois blocs
    sl = nouvelle(prs); n += 1; entete(sl, "Trois idées clés", "Présenter trois axes en parallèle", n)
    for i in range(3):
        x = 0.8 + i * 4.0
        forme(sl, MSO_SHAPE.RECTANGLE, x, 1.9, 3.6, 4.5, fill=BLANC)
        forme(sl, MSO_SHAPE.RECTANGLE, x, 1.9, 3.6, 0.08, fill=ACCENT if i == 1 else PRINCIPAL)
        forme(sl, MSO_SHAPE.OVAL, x + 0.3, 2.3, 0.8, 0.8, fill="E3EBFD", texte=str(i + 1), taille=20,
              couleur=ACCENT, gras=True)
        texte(sl, x + 0.3, 3.3, 3.0, 0.5, f"Titre du bloc {i+1}", 20, PRINCIPAL, True)
        texte(sl, x + 0.3, 3.9, 3.0, 2.3, "Description courte de l'idée en deux à trois lignes. "
              "Rester concret et factuel.", 14, GRIS)

    # 6 Chiffres clés
    sl = nouvelle(prs); n += 1; entete(sl, "Chiffres clés", "Quatre indicateurs à retenir", n)
    chiffres = [("85 %", "de satisfaction"), ("1 200", "utilisateurs actifs"), ("x3", "en deux ans"), ("24 h", "délai moyen")]
    for i, (c, l) in enumerate(chiffres):
        x = 0.8 + i * 3.0
        forme(sl, MSO_SHAPE.RECTANGLE, x, 2.2, 2.7, 3.2, fill=PRINCIPAL if i == 0 else BLANC)
        col = BLANC if i == 0 else ACCENT
        texte(sl, x, 2.6, 2.7, 1.2, c, 48, col, True, PP_ALIGN.CENTER)
        texte(sl, x + 0.2, 3.9, 2.3, 1.2, [l, "Source, période"], 14,
              "C9D7FA" if i == 0 else GRIS, align=PP_ALIGN.CENTER)

    # 7 Frise
    sl = nouvelle(prs); n += 1; entete(sl, "Frise chronologique", "Cinq étapes clés", n)
    trait(sl, 1.2, 3.6, 12.1, 3.6, GRIS, 2)
    for i in range(5):
        cx = 1.6 + i * 2.6
        forme(sl, MSO_SHAPE.OVAL, cx - 0.3, 3.3, 0.6, 0.6, fill=ACCENT if i == 2 else PRINCIPAL)
        texte(sl, cx - 1.2, 2.3, 2.4, 0.8, ["AAAA", "Trimestre"], 14, ACCENT, True, PP_ALIGN.CENTER)
        texte(sl, cx - 1.2, 4.2, 2.4, 0.45, f"Étape {i+1}", 18, PRINCIPAL, True, PP_ALIGN.CENTER)
        texte(sl, cx - 1.2, 4.7, 2.4, 1.4, "Description courte de l'étape.", 13, GRIS, align=PP_ALIGN.CENTER)

    # 8 Comparaison
    sl = nouvelle(prs); n += 1; entete(sl, "Comparaison", "Option A ou option B ?", n)
    for i, (t, c) in enumerate([("Option A", PRINCIPAL), ("Option B", ACCENT)]):
        x = 0.8 + i * 6.0
        forme(sl, MSO_SHAPE.RECTANGLE, x, 1.8, 5.7, 0.8, fill=c, texte=t, taille=22, gras=True)
        forme(sl, MSO_SHAPE.RECTANGLE, x, 2.6, 5.7, 3.9, fill=BLANC)
        texte(sl, x + 0.3, 2.85, 5.1, 3.5, ["+ Avantage principal", "+ Deuxième avantage",
              "– Point de vigilance", "– Contrainte à anticiper", "", "Coût estimé : XX k€", "Délai : X mois"], 16, PRINCIPAL)

    # 9 Citation
    sl = nouvelle(prs, PRINCIPAL); n += 1
    texte(sl, 1.0, 0.3, 2, 2.4, "“", 140, ACCENT, True)
    texte(sl, 1.6, 2.5, 10, 2.2, "Une citation forte, courte et mémorable qui résume l'esprit du projet.",
          32, BLANC, False)
    forme(sl, MSO_SHAPE.RECTANGLE, 1.7, 4.95, 1.0, 0.06, fill=ACCENT)
    texte(sl, 1.6, 5.15, 10, 0.8, ["Prénom Nom", "Fonction, organisation"], 16, "C9D7FA")

    # 10 Organigramme
    sl = nouvelle(prs); n += 1; entete(sl, "Organigramme", "Gouvernance du projet", n)
    forme(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 5.17, 1.7, 3.0, 0.95, fill=PRINCIPAL,
          texte=["Direction", "Prénom Nom"], taille=14, gras=True)
    trait(sl, 6.67, 2.65, 6.67, 3.15); trait(sl, 2.27, 3.15, 11.07, 3.15)
    for i in range(3):
        cx = 2.27 + i * 4.4
        trait(sl, cx, 3.15, cx, 3.5)
        forme(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 1.4, 3.5, 2.8, 0.9, fill=ACCENT,
              texte=[f"Pôle {i+1}", "Responsable"], taille=14, gras=True)
        trait(sl, cx, 4.4, cx, 4.8)
        for j in range(2):
            x = cx - 1.4 + j * 1.45
            if j == 0: trait(sl, cx - 0.7, 4.8, cx + 0.75, 4.8)
            trait(sl, x + 0.68, 4.8, x + 0.68, 5.1)
            forme(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, 5.1, 1.35, 0.75, fill=BLANC, line="C9D7FA",
                  texte=["Équipe", "Nom"], taille=11, couleur=PRINCIPAL)

    # 11 Tableau
    sl = nouvelle(prs); n += 1; entete(sl, "Tableau", "Synthèse des indicateurs", n)
    lignes = [["Indicateur", "Année N-1", "Année N", "Évolution", "Commentaire"]] + \
             [[f"Indicateur {i}", "000", "000", "+ 0 %", "Texte court"] for i in range(1, 6)]
    tb = sl.shapes.add_table(6, 5, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.2)).table
    for r, ligne in enumerate(lignes):
        for c, v in enumerate(ligne):
            cell = tb.cell(r, c); cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(PRINCIPAL if r == 0 else (BLANC if r % 2 else "E9EEF6"))
            cell.text_frame.text = ""
            remplir(cell.text_frame, v, 14, BLANC if r == 0 else PRINCIPAL, r == 0,
                    PP_ALIGN.LEFT if c in (0, 4) else PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

    # 12 Conclusion / contact
    sl = nouvelle(prs, PRINCIPAL); n += 1
    forme(sl, MSO_SHAPE.RECTANGLE, 0, 0, 0.35, 7.5, fill=ACCENT)
    texte(sl, 1.0, 0.9, 7, 0.9, "Merci", 48, BLANC, True)
    texte(sl, 1.0, 1.9, 6.5, 2.6, ["À retenir :", "• Message clé n°1", "• Message clé n°2", "• Prochaine étape et date"],
          20, "C9D7FA")
    forme(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 8.2, 1.5, 4.4, 4.2, fill="2A3758")
    texte(sl, 8.5, 1.8, 3.8, 0.5, "Contact", 22, BLANC, True)
    texte(sl, 8.5, 2.5, 3.8, 3.0, ["Prénom Nom", "Fonction", "", "prenom.nom@exemple.fr", "+33 0 00 00 00 00",
          "www.exemple.fr"], 16, "C9D7FA")

    prs.save(ICI / "modele_presentation.pptx")
    return n


# ===================== KIT SCHEMAS =====================
def kit():
    prs = nouvelle_prs(); n = 0

    def diapo(t, s):
        nonlocal n
        sl = nouvelle(prs); n += 1; entete(sl, t, s, n); return sl

    # Matrice 2x2
    sl = diapo("Matrice 2 × 2", "Positionner des éléments selon deux critères")
    labs = ["Priorité 1\nImportant et urgent", "Priorité 2\nImportant, non urgent",
            "Priorité 3\nUrgent, peu important", "Priorité 4\nNi l'un ni l'autre"]
    cols = [ACCENT, PRINCIPAL, "9DB8F5", "C9D7FA"]
    for i in range(4):
        x, y = 3.2 + (i % 2) * 3.55, 1.6 + (i // 2) * 2.55
        forme(sl, MSO_SHAPE.RECTANGLE, x, y, 3.45, 2.45, fill=cols[i], texte=labs[i].split("\n"),
              taille=16, couleur=BLANC if i < 2 else PRINCIPAL, gras=False)
    texte(sl, 3.2, 6.7, 7, 0.4, "Axe horizontal : critère X  →", 13, GRIS, align=PP_ALIGN.CENTER)
    texte(sl, 0.8, 3.6, 2.3, 0.8, "Axe vertical : critère Y  ↑", 13, GRIS, align=PP_ALIGN.RIGHT)

    # Entonnoir
    sl = diapo("Entonnoir", "Quatre niveaux de conversion")
    for i, (t, v) in enumerate([("Visiteurs", "10 000"), ("Prospects", "2 500"), ("Clients", "500"), ("Fidèles", "150")]):
        w = 8.0 - i * 1.7; x = 6.67 - w / 2; y = 1.7 + i * 1.25
        s = forme(sl, MSO_SHAPE.TRAPEZOID, x, y, w, 1.15, fill=NUANCES[i], texte=f"{t}  –  {v}",
                  taille=18, couleur=BLANC if i < 2 else PRINCIPAL, gras=True)
        s.rotation = 180
        # texte non retourné : on le place dans une zone séparée
        s.text_frame.text = ""
        texte(sl, x, y, w, 1.15, f"{t}  –  {v}", 18, BLANC if i < 2 else PRINCIPAL, True,
              PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

    # Cycle
    sl = diapo("Cycle", "Quatre étapes qui se répètent")
    import math
    cx, cy, R = 6.67, 4.1, 1.9
    forme(sl, MSO_SHAPE.OVAL, cx - R, cy - R, 2 * R, 2 * R, line="C9D7FA")
    noms = ["Planifier", "Réaliser", "Vérifier", "Ajuster"]
    for i, t in enumerate(noms):
        a = math.radians(-90 + i * 90)
        x, y = cx + R * math.cos(a), cy + R * math.sin(a)
        forme(sl, MSO_SHAPE.OVAL, x - 0.8, y - 0.8, 1.6, 1.6, fill=NUANCES[i], texte=t, taille=15,
              couleur=BLANC if i < 3 else PRINCIPAL, gras=True)
    forme(sl, MSO_SHAPE.CIRCULAR_ARROW, cx - 0.9, cy - 0.9, 1.8, 1.8, fill=ACCENT)

    # Pyramide
    sl = diapo("Pyramide", "Quatre niveaux hiérarchiques")
    for i, t in enumerate(["Vision", "Stratégie", "Tactique", "Opérations"]):
        y = 1.6 + i * 1.25; wb = 1.6 * (i + 1)
        if i == 0:
            forme(sl, MSO_SHAPE.ISOSCELES_TRIANGLE, 6.67 - 0.8, y, 1.6, 1.2, fill=NUANCES[0])
            trait(sl, 7.6, y + 0.75, 8.6, y + 0.75, GRIS, 1)
            texte(sl, 8.7, y + 0.5, 3.8, 0.5, "Vision – description courte", 14, PRINCIPAL, True)
        else:
            s = forme(sl, MSO_SHAPE.TRAPEZOID, 6.67 - wb / 2, y, wb, 1.2, fill=NUANCES[i],
                      texte=t, taille=16, couleur=BLANC if i < 3 else PRINCIPAL, gras=True)
            s.adjustments[0] = 0.8 / 1.2  # retrait latéral de 0,8 po de chaque côté

    # Chevrons
    sl = diapo("Processus en chevrons", "Cinq étapes séquentielles")
    for i in range(5):
        x = 0.8 + i * 2.38
        forme(sl, MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON, x, 2.4, 2.5, 1.2,
              fill=NUANCES[i] if i < 3 else ("9DB8F5" if i == 3 else "C9D7FA"),
              texte=f"Étape {i+1}", taille=16, couleur=BLANC if i < 3 else PRINCIPAL, gras=True)
        texte(sl, x + 0.2, 3.9, 2.1, 2, "Description courte de l'étape.", 13, GRIS, align=PP_ALIGN.CENTER)

    # Roadmap
    sl = diapo("Roadmap", "Jalons sur l'année")
    trimestres = ["T1", "T2", "T3", "T4"]
    for i, t in enumerate(trimestres):
        forme(sl, MSO_SHAPE.RECTANGLE, 0.8 + i * 2.93, 1.7, 2.88, 0.5, fill=PRINCIPAL, texte=t, taille=14, gras=True)
    for r, (nom, d, f) in enumerate([("Chantier A", 0, 1.5), ("Chantier B", 1, 3), ("Chantier C", 2.2, 4)]):
        y = 2.6 + r * 1.1
        texte(sl, 0.8, y - 0.35, 3, 0.35, nom, 13, GRIS, True)
        forme(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 0.8 + d * 2.93, y, (f - d) * 2.93, 0.45,
              fill=NUANCES[r + 1], texte="Livrable", taille=12, couleur=BLANC if r < 2 else PRINCIPAL)
    trait(sl, 0.8, 6.1, 12.5, 6.1, GRIS, 2)
    for i, x in enumerate([2.5, 6.2, 10.4]):
        forme(sl, MSO_SHAPE.DIAMOND, x - 0.2, 5.9, 0.4, 0.4, fill=ACCENT)
        texte(sl, x - 1.0, 6.35, 2.0, 0.4, f"Jalon {i+1}", 12, PRINCIPAL, True, PP_ALIGN.CENTER)

    # Venn
    sl = diapo("Diagramme de Venn", "Trois ensembles et leur intersection")
    for (x, y, c, t, tx, ty) in [(4.4, 1.6, PRINCIPAL, "Ensemble A", 4.9, 2.6), (6.6, 1.6, ACCENT, "Ensemble B", 7.6, 2.6),
                                 (5.5, 3.4, "5B8DEF", "Ensemble C", 6.17, 5.3)]:
        s = forme(sl, MSO_SHAPE.OVAL, x, y, 3.0, 3.0, fill=c)
        s.fill._xPr.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr').append(
            __import__("lxml.etree", fromlist=["x"]).fromstring(
                '<a:alpha xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="70000"/>'))
        texte(sl, tx, ty, 1.6, 0.5, t, 14, BLANC, True, PP_ALIGN.CENTER)
    texte(sl, 6.1, 3.75, 1.2, 0.5, "Cœur", 14, BLANC, True, PP_ALIGN.CENTER)

    # Arbre de décision
    sl = diapo("Arbre de décision", "Une question, deux branches, quatre issues")
    forme(sl, MSO_SHAPE.DIAMOND, 5.42, 1.5, 2.5, 1.4, fill=PRINCIPAL, texte="Question ?", taille=15, gras=True)
    for i, (rep, cx) in enumerate([("Oui", 3.67), ("Non", 9.67)]):
        trait(sl, 6.67, 2.9, cx, 3.5, GRIS, 1.5, True)
        texte(sl, (6.67 + cx) / 2 - 0.5, 2.85, 1.0, 0.4, rep, 13, ACCENT, True, PP_ALIGN.CENTER)
        forme(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 1.3, 3.5, 2.6, 0.9, fill=ACCENT, texte=f"Sous-question {i+1}",
              taille=14, gras=True)
        for j, dx in enumerate([-1.5, 1.5]):
            trait(sl, cx, 4.4, cx + dx, 5.1, GRIS, 1.5, True)
            forme(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx + dx - 1.2, 5.1, 2.4, 0.9, fill=BLANC, line=ACCENT,
                  texte=f"Issue {i*2+j+1}", taille=14, couleur=PRINCIPAL)

    # SWOT
    sl = diapo("SWOT", "Forces, faiblesses, opportunités, menaces")
    for i, (t, c) in enumerate([("Forces", PRINCIPAL), ("Faiblesses", ACCENT), ("Opportunités", "5B8DEF"), ("Menaces", GRIS)]):
        x, y = 0.8 + (i % 2) * 5.95, 1.6 + (i // 2) * 2.65
        forme(sl, MSO_SHAPE.RECTANGLE, x, y, 5.8, 0.55, fill=c, texte=t, taille=18, gras=True)
        forme(sl, MSO_SHAPE.RECTANGLE, x, y + 0.55, 5.8, 1.95, fill=BLANC)
        texte(sl, x + 0.2, y + 0.7, 5.4, 1.7, ["• Point 1", "• Point 2", "• Point 3"], 14, PRINCIPAL)

    # Avant / après
    sl = diapo("Avant / Après", "Montrer la transformation")
    for i, (t, c) in enumerate([("Avant", GRIS), ("Après", ACCENT)]):
        x = 0.8 + i * 6.75
        forme(sl, MSO_SHAPE.RECTANGLE, x, 1.7, 5.0, 0.7, fill=c, texte=t, taille=22, gras=True)
        forme(sl, MSO_SHAPE.RECTANGLE, x, 2.4, 5.0, 4.0, fill=BLANC)
        texte(sl, x + 0.3, 2.6, 4.4, 3.6, ["• Situation 1", "• Situation 2", "• Situation 3", "• Situation 4"], 16, PRINCIPAL)
    forme(sl, MSO_SHAPE.RIGHT_ARROW, 6.0, 3.7, 1.3, 0.9, fill=PRINCIPAL)

    prs.save(ICI / "kit_schemas.pptx")
    return n


if __name__ == "__main__":
    print("modele_presentation.pptx :", modele(), "diapos")
    print("kit_schemas.pptx :", kit(), "diapos")
