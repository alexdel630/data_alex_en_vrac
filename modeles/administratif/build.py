# -*- coding: utf-8 -*-
"""Génère les modèles administratifs (devis, facture, suivi des temps, notes de frais).
Données fictives uniquement. Usage : python build.py"""
from pathlib import Path
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule

OUT = Path(__file__).resolve().parent
NUIT, ACCENT = "1F2A44", "2F6FED"
F = "Calibri"
T_FONT = Font(name=F, size=18, bold=True, color=NUIT)
H_FONT = Font(name=F, size=11, bold=True, color="FFFFFF")
B_FONT = Font(name=F, size=11)
BOLD = Font(name=F, size=11, bold=True, color=NUIT)
SMALL = Font(name=F, size=8, color="555555")
H_FILL = PatternFill("solid", fgColor=NUIT)
A_FILL = PatternFill("solid", fgColor=ACCENT)
IN_FILL = PatternFill("solid", fgColor="EEF3FD")
thin = Side(style="thin", color="B7C0D0")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
EUR = '#,##0.00 "€"'
PCT = "0.0%"
DATE = "DD/MM/YYYY"
TVA_RATES = [0.20, 0.10, 0.055, 0.0]


def name(wb, nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def header(ws, row, labels, col=1):
    for i, l in enumerate(labels):
        c = ws.cell(row, col + i, l)
        c.font, c.fill, c.border = H_FONT, H_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def widths(ws, w):
    for k, v in w.items():
        ws.column_dimensions[k].width = v


def print_setup(ws, area, landscape=False):
    ws.print_area = area
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.oddFooter.center.text = "Page &P / &N"
    ws.oddFooter.center.font = F


def list_dv(ws, formula, rng):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error, dv.errorTitle = "Choisir une valeur dans la liste.", "Valeur invalide"
    ws.add_data_validation(dv)
    dv.add(rng)


def params_company(wb, extra=()):
    """Onglet Paramètres : infos société + taux de TVA. Retourne la feuille."""
    ws = wb.create_sheet("Paramètres")
    ws["A1"] = "Paramètres"; ws["A1"].font = T_FONT
    ws["A2"] = "Cellules bleu clair = à compléter. Données fictives à remplacer."
    ws["A2"].font = SMALL
    rows = [
        ("Raison sociale", "Votre Société SAS", "Soc_Nom"),
        ("Adresse", "1 rue Exemple, 75000 Ville", "Soc_Adresse"),
        ("SIRET", "000 000 000 00000 (fictif)", "Soc_SIRET"),
        ("N° TVA intracommunautaire", "FR00 000000000 (fictif)", "Soc_TVA"),
        ("RCS / RM", "RCS Ville 000 000 000 (fictif)", "Soc_RCS"),
        ("Capital social", "1 000 €", "Soc_Capital"),
        ("Téléphone", "01 00 00 00 00", "Soc_Tel"),
        ("E-mail", "contact@votre-societe.example", "Soc_Mail"),
        ("IBAN", "FR76 0000 0000 0000 0000 0000 000 (fictif)", "Soc_IBAN"),
    ] + list(extra)
    r = 4
    for lab, val, nm in rows:
        ws.cell(r, 1, lab).font = BOLD
        c = ws.cell(r, 2, val); c.fill, c.border, c.font = IN_FILL, BOX, B_FONT
        name(wb, nm, f"'Paramètres'!$B${r}")
        r += 1
    r += 1
    ws.cell(r, 1, "Taux de TVA").font = BOLD
    header(ws, r, ["Taux"], col=2)
    first = r + 1
    for i, t in enumerate(TVA_RATES):
        c = ws.cell(first + i, 2, t); c.number_format = PCT; c.fill, c.border = IN_FILL, BOX
    name(wb, "Liste_TVA", f"'Paramètres'!$B${first}:$B${first + len(TVA_RATES) - 1}")
    widths(ws, {"A": 30, "B": 45})
    print_setup(ws, f"A1:B{first + len(TVA_RATES)}")
    return ws, first + len(TVA_RATES) + 1


def piece(kind):
    """Devis ou facture."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Devis" if kind == "devis" else "Facture"
    extra = [("Validité du devis (jours)", 30, "Validite_Jours")] if kind == "devis" else [
        ("Délai de paiement (jours)", 30, "Delai_Paiement"),
        ("Taux pénalités de retard (annuel)", "3 fois le taux d'intérêt légal", "Penalites"),
        ("Escompte pour paiement anticipé", "Néant", "Escompte"),
    ]
    p, _ = params_company(wb, extra)
    widths(ws, {"A": 40, "B": 9, "C": 9, "D": 13, "E": 10, "F": 9, "G": 15})
    ws["A1"] = "=Soc_Nom"; ws["A1"].font = T_FONT
    ws["A2"] = "=Soc_Adresse"
    ws["A3"] = '="SIRET : "&Soc_SIRET&" — TVA : "&Soc_TVA'
    ws["A4"] = '="Tél. : "&Soc_Tel&" — "&Soc_Mail'
    for a in ("A2", "A3", "A4"):
        ws[a].font = SMALL
    ws["E1"] = "DEVIS" if kind == "devis" else "FACTURE"
    ws["E1"].font = Font(name=F, size=22, bold=True, color=ACCENT)
    meta = [("N°", "DEV-2026-001" if kind == "devis" else "FAC-2026-001"),
            ("Date", date(2026, 1, 15))]
    if kind == "devis":
        meta.append(("Valable jusqu'au", "=F3+Validite_Jours"))
    else:
        meta += [("Échéance", "=F3+Delai_Paiement"), ("Réf. devis", "DEV-2026-001")]
    for i, (l, v) in enumerate(meta):
        ws.cell(2 + i, 5, l).font = BOLD
        c = ws.cell(2 + i, 6, v); c.border = BOX
        if not (isinstance(v, str) and v.startswith("=")):
            c.fill = IN_FILL
        if l in ("Date", "Valable jusqu'au", "Échéance"):
            c.number_format = DATE
        ws.merge_cells(start_row=2 + i, start_column=6, end_row=2 + i, end_column=7)
    ws["A7"] = "Client"; ws["A7"].font = BOLD
    for i, v in enumerate(["Client A", "10 avenue Fictive, 00000 Ville", "SIRET client : (fictif)"]):
        c = ws.cell(8 + i, 1, v); c.fill, c.border = IN_FILL, BOX
    ws["A11"] = "Objet :"; ws["A11"].font = BOLD
    ws["B11"] = "Prestation exemple"; ws["B11"].fill = IN_FILL
    ws.merge_cells("B11:G11")

    H, R1, N = 13, 14, 20
    RN = R1 + N - 1
    header(ws, H, ["Désignation", "Quantité", "Unité", "PU HT", "Remise %", "TVA", "Total HT"])
    samples = [("Prestation de conseil", 2, "jour", 650, 0, 0.20),
               ("Fourniture matériel", 3, "u", 120, 0.10, 0.20),
               ("Formation", 1, "forfait", 400, 0, 0.10)]
    for r in range(R1, RN + 1):
        s = samples[r - R1] if r - R1 < len(samples) else (None,) * 6
        for j, v in enumerate(s):
            c = ws.cell(r, 1 + j, v); c.border, c.fill = BOX, IN_FILL
        ws.cell(r, 4).number_format = EUR
        ws.cell(r, 5).number_format = PCT
        ws.cell(r, 6).number_format = PCT
        g = ws.cell(r, 7, f'=IF(OR(B{r}="",D{r}=""),"",ROUND(B{r}*D{r}*(1-N(E{r})),2))')
        g.number_format, g.border = EUR, BOX
    list_dv(ws, "=Liste_TVA", f"F{R1}:F{RN}")
    dvp = DataValidation(type="decimal", operator="between", formula1="0", formula2="1")
    dvp.error = "Remise entre 0 % et 100 %"; ws.add_data_validation(dvp); dvp.add(f"E{R1}:E{RN}")
    dvq = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0")
    ws.add_data_validation(dvq); dvq.add(f"B{R1}:B{RN}"); dvq.add(f"D{R1}:D{RN}")

    r = RN + 2
    ws.cell(r, 5, "Total HT").font = BOLD
    ws.cell(r, 7, f"=SUM(G{R1}:G{RN})").number_format = EUR
    tot_ht = f"G{r}"
    r += 1
    ws.cell(r, 1, "Détail TVA").font = BOLD
    header(ws, r, ["Taux", "Base HT", "Montant TVA"], col=4)
    tva_rows = []
    for i in range(len(TVA_RATES)):
        r += 1
        ws.cell(r, 4, f"=INDEX(Liste_TVA,{i + 1})").number_format = PCT
        ws.cell(r, 5, f"=SUMIFS($G${R1}:$G${RN},$F${R1}:$F${RN},D{r})").number_format = EUR
        ws.cell(r, 6, f"=ROUND(E{r}*D{r},2)").number_format = EUR
        for cc in (4, 5, 6):
            ws.cell(r, cc).border = BOX
        tva_rows.append(r)
    r += 1
    ws.cell(r, 5, "Total TVA").font = BOLD
    ws.cell(r, 7, f"=SUM(F{tva_rows[0]}:F{tva_rows[-1]})").number_format = EUR
    tot_tva = f"G{r}"
    r += 1
    ws.cell(r, 5, "Total TTC").font = H_FONT
    for cc in (5, 6, 7):
        ws.cell(r, cc).fill = A_FILL
    c = ws.cell(r, 7, f"={tot_ht}+{tot_tva}"); c.number_format, c.font = EUR, H_FONT
    ttc = f"G{r}"
    r += 1
    ws.cell(r, 1, f'=IF(ABS(SUM(E{tva_rows[0]}:E{tva_rows[-1]})-{tot_ht})>0.005,'
                  f'"⚠ Des lignes ont un taux de TVA absent de la liste","")').font = Font(name=F, color="C00000")
    if kind == "facture":
        r += 1
        ws.cell(r, 5, "Acompte reçu").font = BOLD
        c = ws.cell(r, 7, 500); c.number_format, c.fill, c.border = EUR, IN_FILL, BOX
        ac = f"G{r}"
        r += 1
        ws.cell(r, 5, "Reste à payer").font = H_FONT
        for cc in (5, 6, 7):
            ws.cell(r, cc).fill = H_FILL
        c = ws.cell(r, 7, f"={ttc}-N({ac})"); c.number_format, c.font = EUR, H_FONT
        r += 2
        lines = [
            '="Conditions de paiement : à "&Delai_Paiement&" jours, échéance le "&TEXT(F4,"dd/mm/yyyy")&", par virement — IBAN : "&Soc_IBAN',
            '="Escompte pour paiement anticipé : "&Escompte',
            '="Pénalités de retard : "&Penalites&", exigibles dès le lendemain de l\'échéance (art. L441-10 C. com.)."',
            "Indemnité forfaitaire pour frais de recouvrement en cas de retard : 40 € (art. D441-5 C. com.).",
            '=Soc_Nom&" — "&Soc_RCS&" — Capital : "&Soc_Capital&" — SIRET : "&Soc_SIRET&" — TVA : "&Soc_TVA',
            "Gabarit indicatif : vérifier les mentions applicables à votre situation (ex. franchise en base : « TVA non applicable, art. 293 B du CGI »).",
        ]
    else:
        r += 2
        lines = [
            '="Devis valable "&Validite_Jours&" jours, jusqu\'au "&TEXT(F4,"dd/mm/yyyy")&"."',
            "Conditions de paiement : 30 % d'acompte à la commande, solde à réception de facture (à adapter).",
            '=Soc_Nom&" — "&Soc_RCS&" — SIRET : "&Soc_SIRET&" — TVA : "&Soc_TVA',
        ]
    for t in lines:
        ws.cell(r, 1, t).font = SMALL
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        r += 1
    if kind == "devis":
        r += 1
        ws.cell(r, 4, "Bon pour accord — date, signature et cachet du client :").font = BOLD
        r += 1
        for rr in range(r, r + 5):
            for cc in range(4, 8):
                ws.cell(rr, cc).border = BOX
        r += 5
    print_setup(ws, f"A1:G{r}")
    ws.freeze_panes = None
    wb.save(OUT / f"{kind}.xlsx")


def suivi_temps():
    wb = Workbook()
    ws = wb.active; ws.title = "Saisie"
    p = wb.create_sheet("Paramètres")
    p["A1"] = "Paramètres"; p["A1"].font = T_FONT
    lists = {"Personnes": ["Personne A", "Personne B", "Personne C"],
             "Projets": ["Projet Alpha", "Projet Beta", "Interne"],
             "Taches": ["Analyse", "Développement", "Réunion", "Gestion"]}
    for j, (nm, vals) in enumerate(lists.items()):
        col = 1 + j
        header(p, 3, [nm], col=col)
        for i in range(30):
            c = p.cell(4 + i, col, vals[i] if i < len(vals) else None); c.fill, c.border = IN_FILL, BOX
        L = p.cell(3, col).column_letter
        name(wb, f"Liste_{nm}", f"'Paramètres'!${L}$4:${L}$33")
        p.column_dimensions[L].width = 22
    p["E3"] = "Année de synthèse"; p["E3"].font = BOLD
    p["F3"] = 2026; p["F3"].fill = IN_FILL
    name(wb, "Annee", "'Paramètres'!$F$3")
    p.column_dimensions["E"].width = 20
    print_setup(p, "A1:F33")

    N = 500
    header(ws, 1, ["Date", "Personne", "Projet", "Tâche", "Heures", "Commentaire", "Mois"])
    widths(ws, {"A": 12, "B": 18, "C": 18, "D": 18, "E": 9, "F": 40, "G": 9})
    import random
    random.seed(1)
    for r in range(2, N + 2):
        if r < 32:
            ws.cell(r, 1, date(2026, 1 + (r - 2) // 3, 1 + (r * 3) % 27))
            ws.cell(r, 2, lists["Personnes"][r % 3]); ws.cell(r, 3, lists["Projets"][r % 3])
            ws.cell(r, 4, lists["Taches"][r % 4]); ws.cell(r, 5, random.choice([1, 2, 3.5, 4, 7]))
            ws.cell(r, 6, "Exemple fictif")
        ws.cell(r, 1).number_format = DATE
        ws.cell(r, 7, f'=IF(A{r}="","",MONTH(A{r}))')
        for cc in range(1, 7):
            ws.cell(r, cc).fill = IN_FILL
    list_dv(ws, "=Liste_Personnes", f"B2:B{N+1}")
    list_dv(ws, "=Liste_Projets", f"C2:C{N+1}")
    list_dv(ws, "=Liste_Taches", f"D2:D{N+1}")
    dvh = DataValidation(type="decimal", operator="between", formula1="0", formula2="24")
    ws.add_data_validation(dvh); dvh.add(f"E2:E{N+1}")
    dvd = DataValidation(type="date", operator="greaterThan", formula1="36526")
    ws.add_data_validation(dvd); dvd.add(f"A2:A{N+1}")
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:G{N+1}"
    print_setup(ws, f"A1:G{N+1}", landscape=True); ws.print_title_rows = "1:1"
    D, Pe, Pr, Ta, Hr = (f"Saisie!${c}$2:${c}${N+1}" for c in "ABCDE")

    s = wb.create_sheet("Synthèse")
    s["A1"] = '="Synthèse des heures — "&Annee'; s["A1"].font = T_FONT
    mois = ["Janv.", "Févr.", "Mars", "Avr.", "Mai", "Juin", "Juil.", "Août", "Sept.", "Oct.", "Nov.", "Déc."]
    s["A3"] = "Heures par projet × mois"; s["A3"].font = BOLD
    header(s, 4, ["Projet"] + mois + ["Total"])
    for i in range(30):
        r = 5 + i
        s.cell(r, 1, f'=IF(INDEX(Liste_Projets,{i+1})="","",INDEX(Liste_Projets,{i+1}))')
        for m in range(12):
            col = s.cell(r, 2 + m).column_letter
            s.cell(r, 2 + m, f'=IF($A{r}="","",SUMIFS({Hr},{Pr},$A{r},{D},">="&DATE(Annee,{m+1},1),{D},"<"&DATE(Annee,{m+2},1)))')
        s.cell(r, 14, f'=IF($A{r}="","",SUM(B{r}:M{r}))').font = BOLD
    r = 35
    s.cell(r, 1, "Total").font = BOLD
    for cc in range(2, 15):
        L = s.cell(r, cc).column_letter
        s.cell(r, cc, f"=SUM({L}5:{L}34)").font = BOLD
    s["A37"] = "Heures par personne"; s["A37"].font = BOLD
    header(s, 38, ["Personne", "Heures (année)", "Part"])
    for i in range(30):
        r = 39 + i
        s.cell(r, 1, f'=IF(INDEX(Liste_Personnes,{i+1})="","",INDEX(Liste_Personnes,{i+1}))')
        s.cell(r, 2, f'=IF($A{r}="","",SUMIFS({Hr},{Pe},$A{r},{D},">="&DATE(Annee,1,1),{D},"<"&DATE(Annee+1,1,1)))')
        c = s.cell(r, 3, f'=IF(OR($A{r}="",SUM($B$39:$B$68)=0),"",B{r}/SUM($B$39:$B$68))'); c.number_format = PCT
    s["A70"] = "Contrôle : heures saisies hors année ou sans projet/personne listé"; s["A70"].font = SMALL
    s["N70"] = f"=SUM({Hr})-N35"; s["N70"].font = SMALL
    widths(s, {"A": 20, "N": 10})
    for cc in range(2, 14):
        s.column_dimensions[s.cell(1, cc).column_letter].width = 8
    s.conditional_formatting.add("B5:M34", FormulaRule(formula=["AND(ISNUMBER(B5),B5>0)"], fill=IN_FILL))
    print_setup(s, "A1:N70", landscape=True)

    x = wb.create_sheet("Croisé")
    x["A1"] = "Tableau croisé : personne × projet (année)"; x["A1"].font = T_FONT
    x["A2"] = "Sans tableau croisé dynamique : formules SUMIFS, 10 premières entrées des listes."; x["A2"].font = SMALL
    x.cell(4, 1, "Personne \\ Projet").font = H_FONT; x.cell(4, 1).fill = H_FILL
    for j in range(10):
        c = x.cell(4, 2 + j, f'=IF(INDEX(Liste_Projets,{j+1})="","",INDEX(Liste_Projets,{j+1}))')
        c.font, c.fill = H_FONT, H_FILL
    x.cell(4, 12, "Total").font = H_FONT; x.cell(4, 12).fill = H_FILL
    for i in range(10):
        r = 5 + i
        x.cell(r, 1, f'=IF(INDEX(Liste_Personnes,{i+1})="","",INDEX(Liste_Personnes,{i+1}))').font = BOLD
        for j in range(10):
            L = x.cell(4, 2 + j).column_letter
            x.cell(r, 2 + j, f'=IF(OR($A{r}="",{L}$4=""),"",SUMIFS({Hr},{Pe},$A{r},{Pr},{L}$4,{D},">="&DATE(Annee,1,1),{D},"<"&DATE(Annee+1,1,1)))')
        x.cell(r, 12, f"=SUM(B{r}:K{r})").font = BOLD
    x.cell(15, 1, "Total").font = BOLD
    for cc in range(2, 13):
        L = x.cell(15, cc).column_letter
        x.cell(15, cc, f"=SUM({L}5:{L}14)").font = BOLD
    widths(x, {"A": 20}); print_setup(x, "A1:L15", landscape=True)
    wb.move_sheet("Paramètres", offset=3)
    wb.save(OUT / "suivi_temps.xlsx")


def notes_de_frais():
    wb = Workbook()
    ws = wb.active; ws.title = "Note de frais"
    p = wb.create_sheet("Paramètres")
    p["A1"] = "Paramètres"; p["A1"].font = T_FONT
    cats = ["Transport", "Hébergement", "Repas", "Indemnités kilométriques", "Fournitures", "Autre"]
    header(p, 3, ["Catégories"])
    for i in range(12):
        c = p.cell(4 + i, 1, cats[i] if i < len(cats) else None); c.fill, c.border = IN_FILL, BOX
    name(wb, "Liste_Categories", "'Paramètres'!$A$4:$A$15")
    header(p, 3, ["Taux TVA"], col=3)
    for i, t in enumerate(TVA_RATES):
        c = p.cell(4 + i, 3, t); c.number_format, c.fill, c.border = PCT, IN_FILL, BOX
    name(wb, "Liste_TVA", f"'Paramètres'!$C$4:$C${3 + len(TVA_RATES)}")
    p["E3"] = "Barème kilométrique (exemple — à mettre à jour avec le barème officiel en vigueur)"; p["E3"].font = BOLD
    p["E4"] = "Puissance du véhicule (CV)"; p["F4"] = 5; p["F4"].fill = IN_FILL
    p["E5"] = "Distance annuelle estimée (km)"; p["F5"] = 6000; p["F5"].fill = IN_FILL
    header(p, 7, ["CV", "≤ 5 000 km : d × a", "5 001–20 000 : d × a + b (a)", "(b)", "> 20 000 : d × a"], col=5)
    bar = [(3, 0.529, 0.316, 1065, 0.370), (4, 0.606, 0.340, 1330, 0.407), (5, 0.636, 0.357, 1395, 0.427),
           (6, 0.665, 0.374, 1457, 0.447), (7, 0.697, 0.394, 1515, 0.470)]
    for i, row in enumerate(bar):
        for j, v in enumerate(row):
            c = p.cell(8 + i, 5 + j, v); c.fill, c.border = IN_FILL, BOX
    p["E14"] = "Taux retenu (€/km, sans la part fixe b)"; p["E14"].font = BOLD
    p["F14"] = ('=IFERROR(IF(F5<=5000,INDEX(F8:F12,MATCH(F4,E8:E12,0)),IF(F5<=20000,'
                'INDEX(G8:G12,MATCH(F4,E8:E12,0)),INDEX(I8:I12,MATCH(F4,E8:E12,0)))),0)')
    p["F14"].number_format = '0.000 "€/km"'
    p["E15"] = "Les valeurs du barème ci-dessus sont des exemples à vérifier (barème URSSAF/impots.gouv de l'année)."
    p["E15"].font = SMALL
    name(wb, "Taux_km", "'Paramètres'!$F$14")
    p["A18"] = "Salarié"; p["B18"] = "Personne A"
    p["A19"] = "Société"; p["B19"] = "Votre Société SAS"
    for a in ("B18", "B19"):
        p[a].fill = IN_FILL
    name(wb, "Salarie", "'Paramètres'!$B$18"); name(wb, "Societe", "'Paramètres'!$B$19")
    widths(p, {"A": 24, "B": 20, "C": 10, "E": 30, "F": 18, "G": 22, "H": 10, "I": 16})
    print_setup(p, "A1:I19", landscape=True)

    ws["A1"] = "Note de frais"; ws["A1"].font = T_FONT
    ws["A2"] = '="Salarié : "&Salarie&" — "&Societe'
    ws["F2"] = "Période :"; ws["G2"] = "Janvier 2026"; ws["G2"].fill = IN_FILL
    H, R1, N = 4, 5, 30
    RN = R1 + N - 1
    header(ws, H, ["Date", "Catégorie", "Description", "Km", "Montant HT", "Taux TVA", "TVA", "TTC", "Justif. (O/N)"])
    widths(ws, {"A": 12, "B": 22, "C": 34, "D": 8, "E": 13, "F": 9, "G": 11, "H": 13, "I": 10})
    ex = [(date(2026, 1, 5), "Transport", "Train aller-retour client A", None, 80, 0.10, "O"),
          (date(2026, 1, 6), "Repas", "Déjeuner client A", None, 30, 0.10, "O"),
          (date(2026, 1, 8), "Indemnités kilométriques", "Visite Client A", 120, None, None, "O")]
    for r in range(R1, RN + 1):
        e = ex[r - R1] if r - R1 < len(ex) else (None,) * 7
        for j, cc in enumerate([1, 2, 3, 4, 5, 6, 9]):
            c = ws.cell(r, cc, e[j]); c.fill, c.border = IN_FILL, BOX
        ws.cell(r, 1).number_format = DATE
        ws.cell(r, 5).number_format = EUR
        ws.cell(r, 6).number_format = PCT
        ws.cell(r, 7, f'=IF(N(D{r})>0,0,IF(E{r}="","",ROUND(E{r}*N(F{r}),2)))').number_format = EUR
        ws.cell(r, 8, f'=IF(N(D{r})>0,ROUND(D{r}*Taux_km,2),IF(E{r}="","",E{r}+G{r}))').number_format = EUR
        for cc in (7, 8):
            ws.cell(r, cc).border = BOX
    list_dv(ws, "=Liste_Categories", f"B{R1}:B{RN}")
    list_dv(ws, "=Liste_TVA", f"F{R1}:F{RN}")
    list_dv(ws, '"O,N"', f"I{R1}:I{RN}")
    ws.conditional_formatting.add(f"A{R1}:I{RN}", FormulaRule(
        formula=[f'AND($B{R1}<>"",$I{R1}<>"O")'], fill=PatternFill("solid", fgColor="FDE2E2")))
    ws.cell(RN + 1, 1, "Km saisis : TTC = km × taux barème (pas de TVA). Ligne rose = justificatif manquant.").font = SMALL
    r = RN + 3
    ws.cell(r, 1, "Totaux par catégorie").font = BOLD
    header(ws, r, ["Catégorie", "", "Km", "HT", "", "TVA", "TTC"], col=2)
    first = r + 1
    for i in range(12):
        r = first + i
        ws.cell(r, 2, f'=IF(INDEX(Liste_Categories,{i+1})="","",INDEX(Liste_Categories,{i+1}))')
        for cc, L in ((4, "D"), (5, "E"), (7, "G"), (8, "H")):
            c = ws.cell(r, cc, f'=IF($B{r}="","",SUMIFS({L}${R1}:{L}${RN},$B${R1}:$B${RN},$B{r}))')
            c.number_format = EUR if cc != 4 else "0"
    r = first + 12
    ws.cell(r, 2, "TOTAL").font = H_FONT
    for cc in range(2, 9):
        ws.cell(r, cc).fill = A_FILL
    for cc, L in ((4, "D"), (5, "E"), (7, "G"), (8, "H")):
        c = ws.cell(r, cc, f"=SUM({L}{R1}:{L}{RN})"); c.font = H_FONT
        c.number_format = EUR if cc != 4 else "0"
    tot = r
    r += 1
    ws.cell(r, 2, f'=IF(ABS(SUM(H{first}:H{first+11})-H{tot})>0.005,"⚠ Catégorie hors liste dans la saisie","")').font = Font(name=F, color="C00000")
    r += 1
    ws.cell(r, 2, "Avance reçue"); c = ws.cell(r, 8, 0); c.fill, c.number_format = IN_FILL, EUR
    r += 1
    ws.cell(r, 2, "Net à rembourser").font = BOLD
    ws.cell(r, 8, f"=H{tot}-H{r-1}").number_format = EUR; ws.cell(r, 8).font = BOLD
    r += 2
    ws.cell(r, 1, "Validation").font = BOLD
    header(ws, r + 1, ["Rôle", "Nom", "Date", "Signature"], col=1)
    ws.merge_cells(start_row=r + 1, start_column=4, end_row=r + 1, end_column=9)
    for i, role in enumerate(["Salarié", "Responsable", "Comptabilité"]):
        rr = r + 2 + i
        ws.cell(rr, 1, role).font = BOLD
        for cc in range(1, 10):
            ws.cell(rr, cc).border = BOX
        ws.cell(rr, 3).number_format = DATE
        ws.row_dimensions[rr].height = 30
    ws.freeze_panes = f"A{R1}"
    print_setup(ws, f"A1:I{r + 4}", landscape=True)
    wb.save(OUT / "notes_de_frais.xlsx")


if __name__ == "__main__":
    piece("devis")
    piece("facture")
    suivi_temps()
    notes_de_frais()
    print("OK :", ", ".join(p.name for p in sorted(OUT.glob("*.xlsx"))))
