# -*- coding: utf-8 -*-
"""Génère les modèles de gestion de projet (Excel + Word).
Usage : python build.py   (fichiers créés à côté du script)
Les formules sont écrites en anglais (openpyxl) ; Excel les affiche en français.
"""
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parent
NAVY, ACCENT = "1F2A44", "2F6FED"
HDR_FILL = PatternFill("solid", fgColor=NAVY)
HDR_FONT = Font(name="Calibri", bold=True, color="FFFFFF")
BASE = Font(name="Calibri", size=11)
TITLE = Font(name="Calibri", size=16, bold=True, color=NAVY)
THIN = Side(style="thin", color="BFC5D2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
ORANGE = PatternFill("solid", fgColor="F8CBAD")
RED = PatternFill("solid", fgColor="FFC7CE")
WRAP = Alignment(wrap_text=True, vertical="top")
TODAY = date(2026, 11, 2)


def header(ws, row, titles, widths=None):
    for c, t in enumerate(titles, 1):
        cell = ws.cell(row=row, column=c, value=t)
        cell.fill, cell.font, cell.border = HDR_FILL, HDR_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = w


def body(ws, r1, r2, c1, c2):
    for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
        for cell in row:
            cell.border, cell.font = BORDER, BASE
            if cell.alignment.wrap_text is not True:
                cell.alignment = WRAP


def dv_list(ws, values, rng, prompt=None):
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    dv.error, dv.errorTitle = "Choisissez une valeur dans la liste.", "Valeur invalide"
    if prompt:
        dv.prompt = prompt
    ws.add_data_validation(dv)
    dv.add(rng)


def guide(wb, title, lines):
    ws = wb.active
    ws.title = "Mode d'emploi"
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws.column_dimensions["A"].width = 110
    for i, l in enumerate(lines, 3):
        c = ws.cell(row=i, column=1, value=l)
        c.font, c.alignment = BASE, Alignment(wrap_text=True)
    ws.cell(row=len(lines) + 4, column=1,
            value="Données d'exemple anonymisées : remplacez-les par les vôtres.").font = \
        Font(name="Calibri", italic=True, color=ACCENT)
    return ws


# ---------------------------------------------------------------- 1. Risques
def registre_risques():
    wb = Workbook()
    guide(wb, "Registre des risques — mode d'emploi", [
        "1. Onglet « Risques » : une ligne par risque (jusqu'à la ligne 200). Les colonnes Catégorie, Probabilité, Impact et Statut ont des listes déroulantes.",
        "2. Probabilité et Impact : 1 (très faible) à 5 (très fort). La criticité (P × I) est calculée automatiquement.",
        "3. Couleurs de criticité : 1-4 vert (faible), 5-9 jaune (modéré), 10-14 orange (élevé), 15-25 rouge (critique).",
        "4. Échéance : passe en rouge si dépassée et risque non clos.",
        "5. Onglet « Matrice » : nombre de risques ouverts (statut ≠ Clos) par case Probabilité × Impact.",
        "6. Onglet « Synthèse » : comptages par niveau, statut et catégorie. Ne modifiez pas les cellules de formule.",
    ])
    ws = wb.create_sheet("Risques")
    cols = ["ID", "Description", "Cause", "Conséquence", "Catégorie", "Propriétaire",
            "Probabilité (1-5)", "Impact (1-5)", "Criticité (P×I)", "Niveau", "Statut",
            "Plan d'action", "Échéance"]
    header(ws, 1, cols, [8, 34, 26, 26, 16, 16, 12, 12, 12, 12, 13, 34, 13])
    data = [
        ("R-001", "Retard de livraison d'un fournisseur clé", "Capacité fournisseur limitée",
         "Décalage du planning", "Fournisseur", "Personne A", 3, 4, "Ouvert",
         "Identifier un fournisseur de secours", TODAY + timedelta(days=20)),
        ("R-002", "Départ d'un expert technique", "Marché tendu", "Perte de compétences",
         "Ressources", "Personne B", 2, 5, "En traitement", "Documenter et former un binôme",
         TODAY + timedelta(days=45)),
        ("R-003", "Dépassement du budget matériel", "Hausse des prix", "Surcoût",
         "Financier", "Personne C", 4, 3, "Ouvert", "Négocier un prix ferme",
         TODAY - timedelta(days=5)),
        ("R-004", "Exigences mal comprises", "Cahier des charges incomplet", "Reprises",
         "Technique", "Personne A", 2, 2, "Clos", "Atelier de validation", TODAY - timedelta(days=30)),
    ]
    for r, d in enumerate(data, 2):
        for c, v in zip([1, 2, 3, 4, 5, 6, 7, 8, 11, 12, 13], d):
            ws.cell(row=r, column=c, value=v)
    N = 200
    for r in range(2, N + 1):
        ws.cell(row=r, column=9, value=f'=IF(AND(ISNUMBER(G{r}),ISNUMBER(H{r})),G{r}*H{r},"")')
        ws.cell(row=r, column=10, value=f'=IF(I{r}="","",IF(I{r}>=15,"Critique",IF(I{r}>=10,"Élevé",IF(I{r}>=5,"Modéré","Faible"))))')
        ws.cell(row=r, column=13).number_format = "DD/MM/YYYY"
        for c in (7, 8, 9, 10):
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="center", vertical="top")
    body(ws, 2, N, 1, 13)
    cats = ["Technique", "Planning", "Financier", "Ressources", "Fournisseur", "Juridique", "Qualité"]
    dv_list(ws, cats, f"E2:E{N}")
    dv_list(ws, ["1", "2", "3", "4", "5"], f"G2:H{N}", "1 = très faible, 5 = très fort")
    dv_list(ws, ["Ouvert", "En traitement", "Clos"], f"K2:K{N}")
    rng = f"I2:J{N}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=["AND($I2<>\"\",$I2>=15)"], fill=RED))
    ws.conditional_formatting.add(rng, FormulaRule(formula=["AND($I2<>\"\",$I2>=10)"], fill=ORANGE))
    ws.conditional_formatting.add(rng, FormulaRule(formula=["AND($I2<>\"\",$I2>=5)"], fill=YELLOW))
    ws.conditional_formatting.add(rng, FormulaRule(formula=["AND($I2<>\"\",$I2>=1)"], fill=GREEN))
    ws.conditional_formatting.add(f"M2:M{N}", FormulaRule(
        formula=['AND($M2<>"",$M2<TODAY(),$K2<>"Clos")'], font=Font(color="9C0006", bold=True), fill=RED))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:M{N}"

    # Matrice
    m = wb.create_sheet("Matrice")
    m["A1"] = "Matrice Probabilité × Impact (risques non clos)"
    m["A1"].font = TITLE
    m["B3"] = "Impact →"
    m["B3"].font = Font(name="Calibri", bold=True)
    m["A4"] = "Probabilité ↓"
    m["A4"].font = Font(name="Calibri", bold=True)
    for i in range(1, 6):
        c = m.cell(row=4, column=2 + i, value=i)
        c.fill, c.font, c.alignment, c.border = HDR_FILL, HDR_FONT, Alignment(horizontal="center"), BORDER
    for k, p in enumerate(range(5, 0, -1)):
        r = 5 + k
        c = m.cell(row=r, column=2, value=p)
        c.fill, c.font, c.alignment, c.border = HDR_FILL, HDR_FONT, Alignment(horizontal="center"), BORDER
        for i in range(1, 6):
            col = m.cell(row=4, column=2 + i).column_letter
            cell = m.cell(row=r, column=2 + i,
                          value=f'=COUNTIFS(Risques!$G$2:$G${N},$B{r},Risques!$H$2:$H${N},{col}$4,Risques!$K$2:$K${N},"<>Clos")')
            s = p * i
            cell.fill = RED if s >= 15 else ORANGE if s >= 10 else YELLOW if s >= 5 else GREEN
            cell.alignment, cell.border = Alignment(horizontal="center", vertical="center"), BORDER
            cell.font = Font(name="Calibri", bold=True, size=14)
        m.row_dimensions[r].height = 38
    m.column_dimensions["A"].width = 14
    for col in "BCDEFG":
        m.column_dimensions[col].width = 11
    m.conditional_formatting.add("C5:G9", CellIsRule(operator="equal", formula=["0"], font=Font(color="A6A6A6")))
    m["A11"] = "Total risques non clos :"
    m["C11"] = "=SUM(C5:G9)"
    m["A11"].font = m["C11"].font = Font(name="Calibri", bold=True)

    # Synthèse
    s = wb.create_sheet("Synthèse")
    s["A1"] = "Synthèse des risques"
    s["A1"].font = TITLE
    header_row = 3
    for c, t in enumerate(["Indicateur", "Valeur"], 1):
        x = s.cell(row=header_row, column=c, value=t)
        x.fill, x.font, x.border = HDR_FILL, HDR_FONT, BORDER
    rows = [
        ("Nombre total de risques", f'=COUNTA(Risques!A2:A{N})'),
        ("Risques non clos", f'=COUNTIFS(Risques!A2:A{N},"<>",Risques!K2:K{N},"<>Clos")'),
        ("Critiques (≥15)", f'=COUNTIF(Risques!J2:J{N},"Critique")'),
        ("Élevés (10-14)", f'=COUNTIF(Risques!J2:J{N},"Élevé")'),
        ("Modérés (5-9)", f'=COUNTIF(Risques!J2:J{N},"Modéré")'),
        ("Faibles (1-4)", f'=COUNTIF(Risques!J2:J{N},"Faible")'),
        ("Criticité moyenne", f'=IFERROR(AVERAGE(Risques!I2:I{N}),0)'),
        ("Échéances dépassées (non clos)", f'=COUNTIFS(Risques!M2:M{N},"<"&TODAY(),Risques!K2:K{N},"<>Clos")'),
        ("Statut : Ouvert", f'=COUNTIF(Risques!K2:K{N},"Ouvert")'),
        ("Statut : En traitement", f'=COUNTIF(Risques!K2:K{N},"En traitement")'),
        ("Statut : Clos", f'=COUNTIF(Risques!K2:K{N},"Clos")'),
    ] + [(f"Catégorie : {c}", f'=COUNTIF(Risques!E2:E{N},"{c}")') for c in cats]
    for i, (a, b) in enumerate(rows, header_row + 1):
        s.cell(row=i, column=1, value=a)
        s.cell(row=i, column=2, value=b)
    body(s, header_row + 1, header_row + len(rows), 1, 2)
    s["B10"].number_format = "0.0"
    s.column_dimensions["A"].width = 34
    s.column_dimensions["B"].width = 12
    s.freeze_panes = "A4"
    wb.save(OUT / "registre_risques.xlsx")


# ---------------------------------------------------------------- 2. RACI
def matrice_raci():
    wb = Workbook()
    guide(wb, "Matrice RACI — mode d'emploi", [
        "R = Réalise, A = Approuve (responsable final, un seul), C = Consulté, I = Informé.",
        "1. Onglet « RACI » : saisissez les tâches en colonne B et les personnes en ligne 1 (colonnes C à H).",
        "2. Dans chaque case, choisissez R, A, C ou I dans la liste déroulante.",
        "3. Colonne « Contrôle » : OK si exactement un A et au moins un R ; sinon le problème est indiqué en rouge.",
        "4. Lignes du bas : charge par personne (nombre de R, A, C, I et total).",
    ])
    ws = wb.create_sheet("RACI")
    people = ["Personne A", "Personne B", "Personne C", "Personne D", "Personne E", "Personne F"]
    header(ws, 1, ["Lot / phase", "Tâche"] + people + ["Nb A", "Nb R", "Contrôle"],
           [16, 36] + [13] * 6 + [8, 8, 26])
    tasks = [
        ("Cadrage", "Rédiger la note de cadrage", ["A", "R", "C", "I", "", ""]),
        ("Cadrage", "Valider le budget", ["A", "C", "", "I", "", ""]),
        ("Conception", "Spécifications fonctionnelles", ["I", "A", "R", "C", "C", ""]),
        ("Conception", "Revue d'architecture", ["", "A", "R", "R", "A", "I"]),
        ("Réalisation", "Développement lot 1", ["I", "", "A", "R", "R", ""]),
        ("Recette", "Plan de recette", ["C", "", "I", "", "", "R"]),
    ]
    N = 41
    for r, (lot, t, v) in enumerate(tasks, 2):
        ws.cell(row=r, column=1, value=lot)
        ws.cell(row=r, column=2, value=t)
        for j, x in enumerate(v):
            if x:
                ws.cell(row=r, column=3 + j, value=x)
    for r in range(2, N + 1):
        ws.cell(row=r, column=9, value=f'=IF(B{r}="","",COUNTIF(C{r}:H{r},"A"))')
        ws.cell(row=r, column=10, value=f'=IF(B{r}="","",COUNTIF(C{r}:H{r},"R"))')
        ws.cell(row=r, column=11, value=(
            f'=IF(B{r}="","",IF(AND(I{r}=1,J{r}>=1),"OK",'
            f'TRIM(IF(I{r}=0,"Aucun A ",IF(I{r}>1,"Plusieurs A ",""))&IF(J{r}=0,"Aucun R",""))))'))
        for c in range(3, 11):
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="center", vertical="top")
    body(ws, 2, N, 1, 11)
    dv_list(ws, ["R", "A", "C", "I"], f"C2:H{N}", "R, A, C ou I")
    ws.conditional_formatting.add(f"K2:K{N}", FormulaRule(formula=['$K2="OK"'], fill=GREEN, font=Font(color="006100")))
    ws.conditional_formatting.add(f"A2:K{N}", FormulaRule(
        formula=['AND($B2<>"",$K2<>"OK")'], fill=RED))
    ws.conditional_formatting.add(f"C2:H{N}", CellIsRule(operator="equal", formula=['"A"'],
                                  font=Font(bold=True, color=ACCENT)))
    ws.freeze_panes = "C2"

    # Charge
    base = N + 2
    lab = ws.cell(row=base, column=2, value="Charge par personne")
    lab.font = Font(name="Calibri", bold=True, color=NAVY, size=12)
    for k, role in enumerate(["R", "A", "C", "I", "Total"]):
        r = base + 1 + k
        c = ws.cell(row=r, column=2, value=f"Nombre de {role}" if role != "Total" else "Total implications")
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BORDER
        for j in range(6):
            col = ws.cell(row=1, column=3 + j).column_letter
            f = (f'=COUNTIF({col}$2:{col}${N},"{role}")' if role != "Total"
                 else f"=SUM({col}{base+1}:{col}{base+4})")
            x = ws.cell(row=r, column=3 + j, value=f)
            x.border, x.alignment = BORDER, Alignment(horizontal="center")
            if role == "Total":
                x.font = Font(name="Calibri", bold=True)
    ws.conditional_formatting.add(f"C{base+5}:H{base+5}", FormulaRule(
        formula=[f"C{base+5}=MAX($C${base+5}:$H${base+5})"], fill=ORANGE))
    wb.save(OUT / "matrice_raci.xlsx")


# ---------------------------------------------------------------- 3. Tableau de bord
def tableau_de_bord():
    wb = Workbook()
    guide(wb, "Tableau de bord projets — mode d'emploi", [
        "1. Onglet « Projets » : une ligne par projet. Météo, Phase et Priorité via listes déroulantes.",
        "2. Météo : Soleil (tout va bien), Nuage (vigilance), Orage (en difficulté).",
        "3. Avancement en % (0 à 100 %), budgets en euros. Écart, % consommé et indicateur sont calculés.",
        "4. Indicateur budget : « Dérive » si le % consommé dépasse l'avancement de plus de 10 points.",
        "5. Onglet « Synthèse » : totaux, répartition par météo et graphique budget prévu / consommé.",
    ])
    ws = wb.create_sheet("Projets")
    cols = ["Projet", "Chef de projet", "Phase", "Priorité", "Météo", "Avancement %",
            "Budget prévu (€)", "Budget consommé (€)", "Écart (€)", "% consommé",
            "Indicateur budget", "Date de fin prévue", "Commentaire"]
    header(ws, 1, cols, [18, 16, 14, 10, 10, 13, 15, 16, 13, 12, 16, 15, 34])
    data = [
        ("Projet A", "Personne A", "Réalisation", "Haute", "Soleil", 0.6, 120000, 65000, date(2027, 3, 31), "Conforme"),
        ("Projet B", "Personne B", "Conception", "Moyenne", "Nuage", 0.3, 80000, 40000, date(2027, 6, 30), "Ressource manquante"),
        ("Projet C", "Personne C", "Recette", "Haute", "Orage", 0.8, 50000, 61000, date(2026, 12, 15), "Surcoût fournisseur"),
        ("Projet D", "Personne A", "Cadrage", "Basse", "Soleil", 0.1, 30000, 2000, date(2027, 9, 30), ""),
    ]
    N = 51
    for r, d in enumerate(data, 2):
        for c, v in zip([1, 2, 3, 4, 5, 6, 7, 8, 12, 13], d):
            ws.cell(row=r, column=c, value=v)
    for r in range(2, N + 1):
        ws.cell(row=r, column=9, value=f'=IF(G{r}="","",G{r}-H{r})')
        ws.cell(row=r, column=10, value=f'=IF(OR(G{r}="",G{r}=0),"",H{r}/G{r})')
        ws.cell(row=r, column=11, value=f'=IF(J{r}="","",IF(H{r}>G{r},"Dépassement",IF(J{r}-F{r}>0.1,"Dérive","Maîtrisé")))')
        ws.cell(row=r, column=6).number_format = "0%"
        ws.cell(row=r, column=10).number_format = "0%"
        for c in (7, 8, 9):
            ws.cell(row=r, column=c).number_format = '#,##0 €;[Red]-#,##0 €'
        ws.cell(row=r, column=12).number_format = "DD/MM/YYYY"
    body(ws, 2, N, 1, 13)
    dv_list(ws, ["Soleil", "Nuage", "Orage"], f"E2:E{N}")
    dv_list(ws, ["Cadrage", "Conception", "Réalisation", "Recette", "Clôturé"], f"C2:C{N}")
    dv_list(ws, ["Haute", "Moyenne", "Basse"], f"D2:D{N}")
    pct = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True)
    pct.error = "Saisissez un pourcentage entre 0 % et 100 %."
    ws.add_data_validation(pct)
    pct.add(f"F2:F{N}")
    cf = ws.conditional_formatting
    cf.add(f"E2:E{N}", CellIsRule(operator="equal", formula=['"Soleil"'], fill=GREEN))
    cf.add(f"E2:E{N}", CellIsRule(operator="equal", formula=['"Nuage"'], fill=YELLOW))
    cf.add(f"E2:E{N}", CellIsRule(operator="equal", formula=['"Orage"'], fill=RED))
    cf.add(f"K2:K{N}", CellIsRule(operator="equal", formula=['"Maîtrisé"'], fill=GREEN))
    cf.add(f"K2:K{N}", CellIsRule(operator="equal", formula=['"Dérive"'], fill=ORANGE))
    cf.add(f"K2:K{N}", CellIsRule(operator="equal", formula=['"Dépassement"'], fill=RED))
    from openpyxl.formatting.rule import DataBarRule
    cf.add(f"F2:F{N}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=ACCENT))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:M{N}"

    s = wb.create_sheet("Synthèse")
    s["A1"] = "Synthèse du portefeuille"
    s["A1"].font = TITLE
    for c, t in enumerate(["Indicateur", "Valeur"], 1):
        x = s.cell(row=3, column=c, value=t)
        x.fill, x.font, x.border = HDR_FILL, HDR_FONT, BORDER
    P = f"Projets!$A$2:$A${N}"
    rows = [
        ("Nombre de projets", f"=COUNTA({P})", "0"),
        ("Budget prévu total", f"=SUM(Projets!G2:G{N})", '#,##0 €'),
        ("Budget consommé total", f"=SUM(Projets!H2:H{N})", '#,##0 €'),
        ("Écart total", "=B5-B6", '#,##0 €;[Red]-#,##0 €'),
        ("% budget consommé", '=IF(B5=0,0,B6/B5)', "0%"),
        ("Avancement moyen", f"=IFERROR(AVERAGE(Projets!F2:F{N}),0)", "0%"),
        ("Avancement pondéré par budget", f"=IF(B5=0,0,SUMPRODUCT(Projets!F2:F{N},Projets!G2:G{N})/B5)", "0%"),
        ("Projets en dépassement", f'=COUNTIF(Projets!K2:K{N},"Dépassement")', "0"),
        ("Météo Soleil", f'=COUNTIF(Projets!E2:E{N},"Soleil")', "0"),
        ("Météo Nuage", f'=COUNTIF(Projets!E2:E{N},"Nuage")', "0"),
        ("Météo Orage", f'=COUNTIF(Projets!E2:E{N},"Orage")', "0"),
    ]
    for i, (a, b, fmt) in enumerate(rows, 4):
        s.cell(row=i, column=1, value=a)
        s.cell(row=i, column=2, value=b).number_format = fmt
    body(s, 4, 3 + len(rows), 1, 2)
    s.column_dimensions["A"].width = 32
    s.column_dimensions["B"].width = 16
    s.freeze_panes = "A4"

    ch = BarChart()
    ch.type = "col"
    ch.title = "Budget prévu vs consommé par projet"
    ch.y_axis.title = "€"
    ch.height, ch.width = 8, 18
    ch.add_data(Reference(ws, min_col=7, max_col=8, min_row=1, max_row=1 + len(data)), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=1 + len(data)))
    ch.series[0].graphicalProperties.solidFill = NAVY
    ch.series[1].graphicalProperties.solidFill = ACCENT
    s.add_chart(ch, "D3")
    wb.save(OUT / "tableau_de_bord_projets.xlsx")


# ---------------------------------------------------------------- 4. Actions
def suivi_actions():
    wb = Workbook()
    guide(wb, "Suivi des actions de réunion — mode d'emploi", [
        "1. Onglet « Réunions » : journal des réunions (code, date, type, animateur, participants).",
        "2. Onglet « Actions » : une ligne par action. La colonne Réunion propose les codes saisis dans le journal.",
        "3. Retard (jours) = AUJOURDHUI() − échéance, si l'action n'est pas terminée ; mis à jour à chaque ouverture.",
        "4. Les actions en retard apparaissent en rouge, celles à échéance sous 7 jours en orange.",
        "5. Utilisez les filtres automatiques (ligne d'en-tête) pour trier par responsable ou statut.",
    ])
    rj = wb.create_sheet("Réunions")
    header(rj, 1, ["Code réunion", "Date", "Type", "Animateur", "Participants", "Objet"], [14, 12, 16, 16, 34, 40])
    reus = [("REU-01", TODAY - timedelta(days=21), "Comité de pilotage", "Personne A", "Personne A, B, C", "Lancement Projet A"),
            ("REU-02", TODAY - timedelta(days=7), "Point hebdo", "Personne B", "Personne B, C, D", "Avancement lot 1")]
    for r, d in enumerate(reus, 2):
        for c, v in enumerate(d, 1):
            rj.cell(row=r, column=c, value=v)
    for r in range(2, 101):
        rj.cell(row=r, column=2).number_format = "DD/MM/YYYY"
    body(rj, 2, 100, 1, 6)
    dv_list(rj, ["Comité de pilotage", "Comité projet", "Point hebdo", "Atelier", "Revue"], "C2:C100")
    rj.freeze_panes = "B2"
    rj.auto_filter.ref = "A1:F100"

    ws = wb.create_sheet("Actions")
    header(ws, 1, ["N°", "Réunion", "Action", "Responsable", "Date de création", "Échéance",
                   "Statut", "Retard (jours)", "Commentaire"], [7, 12, 40, 16, 14, 13, 13, 13, 34])
    acts = [
        (1, "REU-01", "Constituer l'équipe projet", "Personne A", TODAY - timedelta(days=21), TODAY - timedelta(days=10), "Terminé"),
        (2, "REU-01", "Valider le budget", "Personne C", TODAY - timedelta(days=21), TODAY - timedelta(days=3), "En cours"),
        (3, "REU-02", "Mettre à jour le planning", "Personne B", TODAY - timedelta(days=7), TODAY + timedelta(days=4), "À faire"),
        (4, "REU-02", "Préparer la démonstration", "Personne D", TODAY - timedelta(days=7), TODAY + timedelta(days=20), "À faire"),
    ]
    N = 300
    for r, d in enumerate(acts, 2):
        for c, v in enumerate(d, 1):
            ws.cell(row=r, column=c, value=v)
    for r in range(2, N + 1):
        ws.cell(row=r, column=8, value=f'=IF(OR(F{r}="",G{r}="Terminé",G{r}="Annulé"),"",MAX(0,TODAY()-F{r}))')
        for c in (5, 6):
            ws.cell(row=r, column=c).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=8).alignment = Alignment(horizontal="center", vertical="top")
    body(ws, 2, N, 1, 9)
    dv_list(ws, ["À faire", "En cours", "Terminé", "Annulé"], f"G2:G{N}")
    dvr = DataValidation(type="list", formula1="='Réunions'!$A$2:$A$100", allow_blank=True)
    ws.add_data_validation(dvr)
    dvr.add(f"B2:B{N}")
    cf = ws.conditional_formatting
    cf.add(f"A2:I{N}", FormulaRule(formula=['AND($A2<>"",ISNUMBER($H2),$H2>0)'], fill=RED))
    cf.add(f"A2:I{N}", FormulaRule(formula=['AND($F2<>"",$G2<>"Terminé",$G2<>"Annulé",$F2>=TODAY(),$F2-TODAY()<=7)'], fill=ORANGE))
    cf.add(f"A2:I{N}", FormulaRule(formula=['$G2="Terminé"'], font=Font(color="808080")))
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:I{N}"
    wb.save(OUT / "suivi_actions_reunions.xlsx")


# ---------------------------------------------------------------- 5. Word
def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def word_table(doc, headers, rows):
    t = doc.add_table(rows=1 + rows, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, NAVY)
    return t


def compte_rendu():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for name in ("Title", "Heading 1", "Heading 2"):
        s = doc.styles[name]
        s.font.name = "Calibri"
        s.font.color.rgb = RGBColor.from_string(NAVY if name != "Heading 2" else ACCENT)
    hp = doc.sections[0].header.paragraphs[0]
    hp.text = "[Nom du projet] — Compte rendu de réunion"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp = doc.sections[0].footer.paragraphs[0]
    fp.text = "Document interne — diffusion : participants"
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Compte rendu de réunion", 0)
    t = doc.add_table(rows=5, cols=2)
    t.style = "Table Grid"
    for i, (k, v) in enumerate([("Projet", "Projet A"), ("Objet", "[Objet de la réunion]"),
                                ("Date et heure", "JJ/MM/AAAA — HHhMM"), ("Lieu / visio", "[Salle ou lien]"),
                                ("Rédacteur", "Personne A")]):
        t.cell(i, 0).text = k
        t.cell(i, 0).paragraphs[0].runs[0].bold = True
        shade(t.cell(i, 0), "E8ECF5")
        t.cell(i, 1).text = v

    doc.add_heading("1. Participants", 1)
    p = word_table(doc, ["Nom", "Fonction / entité", "Présent", "Excusé", "Diffusion"], 3)
    for i, n in enumerate(["Personne A", "Personne B", "Personne C"], 1):
        p.cell(i, 0).text = n

    doc.add_heading("2. Ordre du jour", 1)
    for x in ["Validation du compte rendu précédent", "Point d'avancement", "Risques et points bloquants", "Questions diverses"]:
        doc.add_paragraph(x, style="List Number")

    doc.add_heading("3. Points abordés", 1)
    doc.add_heading("3.1 [Sujet]", 2)
    doc.add_paragraph("[Synthèse des échanges]")

    doc.add_heading("4. Décisions", 1)
    d = word_table(doc, ["N°", "Décision", "Validée par"], 2)
    d.cell(1, 0).text, d.cell(2, 0).text = "D1", "D2"

    doc.add_heading("5. Actions", 1)
    a = word_table(doc, ["N°", "Action", "Responsable", "Échéance", "Statut"], 3)
    for i in range(1, 4):
        a.cell(i, 0).text = f"A{i}"
        a.cell(i, 4).text = "À faire"
    doc.add_paragraph("Reporter ces actions dans suivi_actions_reunions.xlsx.").runs[0].italic = True

    doc.add_heading("6. Prochaine réunion", 1)
    for k in ["Date et heure : JJ/MM/AAAA — HHhMM", "Lieu / visio : ", "Ordre du jour prévisionnel : "]:
        doc.add_paragraph(k, style="List Bullet")
    doc.save(OUT / "compte_rendu_reunion.docx")


if __name__ == "__main__":
    registre_risques()
    matrice_raci()
    tableau_de_bord()
    suivi_actions()
    compte_rendu()
    print("Fichiers générés dans", OUT)
