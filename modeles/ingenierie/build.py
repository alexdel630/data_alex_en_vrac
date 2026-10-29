# -*- coding: utf-8 -*-
"""Génère les modèles d'ingénierie : CDC, traçabilité, AMDEC, revue de conception.
Usage : python build.py   (sorties dans le même dossier)"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent
NAVY, ACCENT, FONT = "1F2A44", "2F6FED", "Calibri"
GREY = RGBColor(0x80, 0x80, 0x80)

# ----------------------------------------------------------------- Word
def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcol)
    tcPr.append(shd)

def new_doc():
    d = Document()
    for s in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Title"):
        st = d.styles[s]; st.font.name = FONT
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    d.styles["Normal"].font.size = Pt(11)
    for s, sz in (("Heading 1", 16), ("Heading 2", 13)):
        d.styles[s].font.size = Pt(sz); d.styles[s].font.color.rgb = RGBColor.from_string(NAVY)
    d.styles["Title"].font.color.rgb = RGBColor.from_string(NAVY)
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)
    return d

def aide(d, txt):
    r = d.add_paragraph().add_run(f"[{txt}]")
    r.italic = True; r.font.color.rgb = GREY

def table(d, headers, rows, widths=None):
    t = d.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, NAVY)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(v)
            if v.startswith("["): r.italic = True; r.font.color.rgb = GREY
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    d.add_paragraph()
    return t

def toc(d):
    p = d.add_paragraph()
    for kind, text in (("begin", None), (None, 'TOC \\o "1-3" \\h \\z \\u'), ("separate", None), ("end", None)):
        r = p.add_run()
        if kind:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), kind); r._r.append(f)
            if kind == "separate":
                p.add_run("[Clic droit > Mettre à jour les champs pour générer le sommaire]").italic = True
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._r.append(it)

def cover(d, titre, sous_titre, champs):
    for _ in range(6): d.add_paragraph()
    p = d.add_paragraph(titre, style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(sous_titre); r.font.size = Pt(14); r.font.color.rgb = RGBColor.from_string(ACCENT)
    d.add_paragraph()
    table(d, ["Champ", "Valeur"], [[k, v] for k, v in champs], [5, 11])
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def build_cdc():
    d = new_doc()
    cover(d, "CAHIER DES CHARGES", "[Nom du projet / produit]", [
        ("Référence", "[CDC-XXX-001]"), ("Client / MOA", "[Société A]"), ("Rédacteur", "[Prénom Nom]"),
        ("Version", "[0.1]"), ("Date", "[JJ/MM/AAAA]"), ("Diffusion", "[Interne / Confidentiel]")])
    d.add_heading("Historique des versions", 1)
    table(d, ["Version", "Date", "Auteur", "Modifications", "Validé par"],
          [["0.1", "[JJ/MM/AAAA]", "[Auteur]", "[Création]", "[Nom]"], ["", "", "", "", ""]])
    d.add_heading("Sommaire", 1); toc(d)
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    d.add_heading("1. Contexte", 1)
    aide(d, "Décrire la situation actuelle, le besoin à l'origine du projet et les enjeux (2 à 5 paragraphes).")
    d.add_heading("2. Objectifs", 1)
    aide(d, "Objectifs mesurables (SMART). Ex. : réduire le temps de traitement de 30 % d'ici T4.")
    table(d, ["N°", "Objectif", "Indicateur", "Cible"], [["OBJ-01", "[Objectif]", "[KPI]", "[Valeur]"]])
    d.add_heading("3. Périmètre", 1)
    d.add_heading("3.1 Inclus", 2); aide(d, "Fonctions, sites, populations couverts.")
    d.add_heading("3.2 Exclus", 2); aide(d, "Ce qui est explicitement hors périmètre.")
    d.add_heading("4. Parties prenantes", 1)
    table(d, ["Rôle", "Nom / entité", "Responsabilités", "Contact"],
          [["MOA", "[Entité]", "[Expression du besoin, recette]", "[email]"],
           ["MOE", "[Entité]", "[Conception, réalisation]", "[email]"]])
    d.add_heading("5. Exigences fonctionnelles", 1)
    aide(d, "Priorité MoSCoW : M = Must (indispensable), S = Should, C = Could, W = Won't (pas cette version).")
    table(d, ["ID", "Exigence", "Priorité (MoSCoW)", "Critère d'acceptation"],
          [["EF-001", "[Le système doit permettre de …]", "M", "[Test mesurable]"],
           ["EF-002", "[…]", "S", "[…]"]], [2, 7, 3, 5])
    d.add_heading("6. Exigences non fonctionnelles", 1)
    table(d, ["ID", "Catégorie", "Exigence", "Critère"],
          [["ENF-001", "Performance", "[Temps de réponse < 2 s]", "[95e percentile]"],
           ["ENF-002", "Sécurité", "[…]", "[…]"], ["ENF-003", "Disponibilité", "[…]", "[…]"]])
    d.add_heading("7. Contraintes", 1)
    aide(d, "Techniques, réglementaires, budgétaires, calendaires, normes applicables.")
    d.add_heading("8. Interfaces", 1)
    table(d, ["ID", "Système externe", "Flux / protocole", "Sens", "Fréquence"],
          [["IF-01", "[Système B]", "[API REST / fichier]", "[Entrant]", "[Quotidien]"]])
    d.add_heading("9. Livrables", 1)
    table(d, ["Livrable", "Format", "Date prévue", "Responsable"], [["[Dossier de conception]", "[PDF]", "[JJ/MM]", "[Nom]"]])
    d.add_heading("10. Planning", 1)
    table(d, ["Jalon", "Description", "Date", "Critère de passage"],
          [["J0", "Lancement", "[JJ/MM]", "[CDC validé]"], ["J1", "Revue de conception", "[JJ/MM]", "[Go]"],
           ["J2", "Recette", "[JJ/MM]", "[PV signé]"]])
    d.add_heading("11. Critères de recette", 1)
    aide(d, "Conditions de réception : % d'exigences M validées, absence d'anomalie bloquante, documentation remise…")
    d.add_heading("12. Glossaire", 1)
    table(d, ["Terme", "Définition"], [["MoSCoW", "Méthode de priorisation Must/Should/Could/Won't"], ["[Terme]", "[Définition]"]])
    d.save(OUT / "cahier_des_charges.docx")

def build_revue():
    d = new_doc()
    cover(d, "COMPTE RENDU DE REVUE DE CONCEPTION", "[Projet] – [Revue préliminaire / détaillée]", [
        ("Référence", "[RC-XXX-001]"), ("Date de la revue", "[JJ/MM/AAAA]"), ("Animateur", "[Prénom Nom]"),
        ("Version du dossier revu", "[1.0]")])
    d.add_heading("1. Participants", 1)
    table(d, ["Nom", "Fonction", "Présent / Excusé"], [["[Prénom Nom]", "[Chef de projet]", "[Présent]"]])
    d.add_heading("2. Documents revus", 1)
    table(d, ["Document", "Référence", "Version"], [["[Dossier de conception]", "[REF]", "[1.0]"]])
    d.add_heading("3. Checklist de revue", 1)
    aide(d, "Statut : OK / NOK / N/A. Tout NOK doit générer un point ouvert.")
    items = ["Toutes les exigences du CDC sont couvertes", "Architecture justifiée (alternatives étudiées)",
             "Interfaces définies et validées", "Analyse de risques / AMDEC à jour",
             "Exigences non fonctionnelles adressées", "Plan de vérification / tests défini",
             "Conformité normes et réglementation", "Maintenabilité et évolutivité prises en compte",
             "Coûts et planning cohérents", "Documentation complète et à jour"]
    table(d, ["N°", "Point de contrôle", "Statut", "Commentaire"],
          [[f"{i:02d}", it, "[OK/NOK/N/A]", ""] for i, it in enumerate(items, 1)], [1.2, 8, 2.5, 5])
    d.add_heading("4. Points ouverts", 1)
    table(d, ["ID", "Description", "Gravité", "Responsable", "Échéance", "Statut"],
          [["PO-01", "[…]", "[Bloquant/Majeur/Mineur]", "[Nom]", "[JJ/MM]", "[Ouvert]"]])
    d.add_heading("5. Décision", 1)
    for opt in ("GO – passage à la phase suivante", "GO sous réserves – levée des points ouverts listés",
                "NO-GO – nouvelle revue nécessaire"):
        d.add_paragraph(f"☐  {opt}")
    aide(d, "Justification de la décision et conditions de levée des réserves.")
    d.add_heading("6. Signatures", 1)
    table(d, ["Rôle", "Nom", "Date", "Signature"], [["Animateur", "", "", ""], ["Responsable technique", "", "", ""], ["MOA", "", "", ""]])
    d.save(OUT / "revue_de_conception.docx")

# ----------------------------------------------------------------- Excel
HFILL = PatternFill("solid", fgColor=NAVY)
HFONT = Font(name=FONT, bold=True, color="FFFFFF")
BFONT = Font(name=FONT)
THIN = Side(style="thin", color="BFBFBF"); BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILLS = {k: PatternFill("solid", fgColor=v) for k, v in
         dict(green="C6EFCE", orange="FFEB9C", red="FFC7CE", blue="DCE6FB").items()}

def header(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row, i, h); c.fill = HFILL; c.font = HFONT; c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 32

def body(ws, r1, r2, ncol):
    for r in range(r1, r2 + 1):
        for c in range(1, ncol + 1):
            cell = ws.cell(r, c); cell.border = BORDER
            cell.font = BFONT; cell.alignment = Alignment(vertical="top", wrap_text=True)

def title(ws, txt):
    ws["A1"] = txt; ws["A1"].font = Font(name=FONT, bold=True, size=14, color=NAVY)

def dv_list(ws, formula, rng):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True); dv.add(rng); ws.add_data_validation(dv)

def mode_emploi(wb, lignes):
    ws = wb.create_sheet("Mode d'emploi", 0); title(ws, "Mode d'emploi"); ws.column_dimensions["A"].width = 110
    for i, l in enumerate(lignes, 3):
        ws.cell(i, 1, l).font = Font(name=FONT, bold=l.endswith(":"), color=NAVY if l.endswith(":") else "000000")
        ws.cell(i, 1).alignment = Alignment(wrap_text=True)

def build_trace():
    wb = Workbook(); ws = wb.active; ws.title = "Traçabilité"; N = 200
    title(ws, "Matrice de traçabilité des exigences")
    hd = ["ID exigence", "Exigence", "Priorité", "Source (CDC §)", "Élément de conception", "ID test",
          "Description du test", "Résultat test", "Statut", "Couvert ?", "Commentaire"]
    header(ws, 3, hd, [12, 40, 10, 12, 28, 10, 32, 13, 13, 10, 28])
    ex = [["EF-001", "[Le système doit …]", "M", "§5", "[Module A]", "T-001", "[Cas nominal]", "Réussi", "Validé"],
          ["EF-002", "[…]", "S", "§5", "[Module B]", "T-002", "[…]", "Échoué", "En cours"],
          ["ENF-001", "[Temps de réponse < 2 s]", "M", "§6", "", "", "", "Non exécuté", "À faire"]]
    last = 3 + N
    for r in range(4, last + 1):
        if r - 4 < len(ex):
            for c, v in enumerate(ex[r - 4], 1): ws.cell(r, c, v)
        ws.cell(r, 10, f'=IF(A{r}="","",IF(AND(E{r}<>"",F{r}<>""),"Oui","Non"))')
    body(ws, 4, last, len(hd))
    dv_list(ws, '"M,S,C,W"', f"C4:C{last}")
    dv_list(ws, '"Non exécuté,Réussi,Échoué,Bloqué"', f"H4:H{last}")
    dv_list(ws, '"À faire,En cours,Validé,Rejeté"', f"I4:I{last}")
    rng = f"A4:K{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$I4="Validé"'], fill=FILLS["green"]))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$I4="Rejeté"'], fill=FILLS["red"]))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$I4="En cours"'], fill=FILLS["orange"]))
    ws.conditional_formatting.add(f"J4:J{last}", CellIsRule(operator="equal", formula=['"Non"'], font=Font(color="C00000", bold=True)))
    ws.freeze_panes = "C4"; ws.auto_filter.ref = f"A3:K{last}"

    s = wb.create_sheet("Synthèse"); title(s, "Synthèse de couverture et de validation")
    T = f"'Traçabilité'!"
    kp = [("Nombre d'exigences", f'=COUNTA({T}A4:A{last})'),
          ("Exigences couvertes (conception + test)", f'=COUNTIF({T}J4:J{last},"Oui")'),
          ("Taux de couverture", "=IF(B3=0,0,B4/B3)"),
          ("Exigences validées", f'=COUNTIF({T}I4:I{last},"Validé")'),
          ("Taux de validation global", "=IF(B3=0,0,B6/B3)")]
    for i, (k, f) in enumerate(kp, 3):
        s.cell(i, 1, k).font = Font(name=FONT, bold=True); s.cell(i, 2, f).font = BFONT
        s.cell(i, 1).border = s.cell(i, 2).border = BORDER
    s["B5"].number_format = s["B7"].number_format = "0.0%"
    header(s, 9, ["Priorité", "Total", "Validées", "% validé", "Couvertes", "% couvert"], [40, 12, 12, 12, 12, 12])
    for i, p in enumerate("MSCW", 10):
        s.cell(i, 1, p)
        s.cell(i, 2, f'=COUNTIFS({T}C4:C{last},A{i},{T}A4:A{last},"<>")')
        s.cell(i, 3, f'=COUNTIFS({T}C4:C{last},A{i},{T}I4:I{last},"Validé")')
        s.cell(i, 4, f'=IF(B{i}=0,0,C{i}/B{i})')
        s.cell(i, 5, f'=COUNTIFS({T}C4:C{last},A{i},{T}J4:J{last},"Oui")')
        s.cell(i, 6, f'=IF(B{i}=0,0,E{i}/B{i})')
        s.cell(i, 4).number_format = s.cell(i, 6).number_format = "0.0%"
    body(s, 10, 13, 6)
    for col in ("D", "F"):
        s.conditional_formatting.add(f"{col}10:{col}13", CellIsRule(operator="greaterThanOrEqual", formula=["1"], fill=FILLS["green"]))
        s.conditional_formatting.add(f"{col}10:{col}13", CellIsRule(operator="lessThan", formula=["0.5"], fill=FILLS["red"]))
    s.freeze_panes = "A3"
    mode_emploi(wb, ["Objet :", "Relier chaque exigence du cahier des charges à sa conception, ses tests et son statut.",
        "Saisie :", "1. Onglet Traçabilité : une ligne par couple exigence/test (lignes 4 à 203).",
        "2. Priorité, Résultat test et Statut se choisissent dans les listes déroulantes.",
        "3. La colonne « Couvert ? » est calculée : Oui si conception ET test renseignés. Ne pas la saisir.",
        "Lecture :", "Onglet Synthèse : taux de couverture et % validé par priorité MoSCoW (COUNTIFS).",
        "Couleurs : vert = validé, orange = en cours, rouge = rejeté.",
        "Attention :", "Si une exigence a plusieurs tests (plusieurs lignes), elle est comptée plusieurs fois dans la synthèse."])
    wb.save(OUT / "tracabilite_exigences.xlsx")

def build_amdec():
    wb = Workbook(); ws = wb.active; ws.title = "AMDEC"; N = 100
    title(ws, "Fiche AMDEC – [Produit / Processus]")
    ws["A2"] = "Seuils : criticité < [Seuils!B3] acceptable ; < [Seuils!B4] à surveiller ; au-delà action obligatoire."
    ws["A2"].font = Font(name=FONT, italic=True, color="808080")
    hd = ["N°", "Élément / Fonction", "Mode de défaillance", "Effet", "Cause", "Détection actuelle",
          "G", "O", "D", "Criticité (C=G×O×D)", "Action corrective", "Responsable", "Échéance",
          "G'", "O'", "D'", "Criticité résiduelle", "Gain"]
    header(ws, 3, hd, [6, 22, 24, 24, 24, 20, 5, 5, 5, 11, 28, 14, 11, 5, 5, 5, 11, 9])
    ex = [[1, "[Pompe]", "[Ne démarre pas]", "[Arrêt ligne]", "[Défaut alimentation]", "[Contrôle visuel]", 8, 4, 5,
           None, "[Ajout capteur de courant]", "[Nom]", "[JJ/MM]", 8, 2, 2],
          [2, "[Joint]", "[Fuite]", "[Perte de fluide]", "[Usure]", "[Inspection mensuelle]", 5, 3, 3,
           None, "", "", "", None, None, None]]
    last = 3 + N
    for r in range(4, last + 1):
        if r - 4 < len(ex):
            for c, v in enumerate(ex[r - 4], 1):
                if v is not None: ws.cell(r, c, v)
        ws.cell(r, 10, f'=IF(COUNT(G{r}:I{r})=3,G{r}*H{r}*I{r},"")')
        ws.cell(r, 17, f'=IF(COUNT(N{r}:P{r})=3,N{r}*O{r}*P{r},"")')
        ws.cell(r, 18, f'=IF(AND(ISNUMBER(J{r}),ISNUMBER(Q{r}),J{r}>0),(J{r}-Q{r})/J{r},"")')
        ws.cell(r, 18).number_format = "0%"
    body(ws, 4, last, len(hd))
    for r in range(4, last + 1):
        for c in (7, 8, 9, 10, 14, 15, 16, 17): ws.cell(r, c).alignment = Alignment(horizontal="center", vertical="top")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="10", allow_blank=True,
                        showErrorMessage=True, errorTitle="Note invalide", error="Saisir un entier de 1 à 10 (voir onglet Barèmes).")
    dv.add(f"G4:I{last}"); dv.add(f"N4:P{last}"); ws.add_data_validation(dv)
    for col in ("J", "Q"):
        rg = f"{col}4:{col}{last}"
        ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND(ISNUMBER({col}4),{col}4>=Seuils!$B$4)'], fill=FILLS["red"], font=Font(bold=True, color="9C0006")))
        ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND(ISNUMBER({col}4),{col}4>=Seuils!$B$3)'], fill=FILLS["orange"]))
        ws.conditional_formatting.add(rg, FormulaRule(formula=[f'ISNUMBER({col}4)'], fill=FILLS["green"]))
    ws.freeze_panes = "C4"; ws.auto_filter.ref = f"A3:R{last}"

    b = wb.create_sheet("Barèmes"); title(b, "Barèmes de cotation (1 à 10)")
    header(b, 3, ["Note", "Gravité (G)", "Occurrence (O)", "Détection (D)"], [7, 45, 40, 45])
    rows = [(1, "Sans effet perceptible", "Quasi impossible (< 1/1 000 000)", "Détection certaine"),
            (2, "Effet très mineur", "Très rare", "Détection quasi certaine"),
            (3, "Effet mineur, gêne légère", "Rare", "Très forte probabilité de détection"),
            (4, "Gêne notable, sans perte de fonction", "Peu fréquent", "Forte probabilité"),
            (5, "Dégradation partielle de fonction", "Occasionnel", "Probabilité modérée"),
            (6, "Perte partielle de fonction", "Assez fréquent", "Faible probabilité"),
            (7, "Perte de fonction principale", "Fréquent", "Très faible probabilité"),
            (8, "Arrêt majeur, client très insatisfait", "Très fréquent", "Détection improbable"),
            (9, "Risque sécurité / non-conformité, avec alerte", "Quasi systématique", "Détection très improbable"),
            (10, "Risque sécurité / réglementaire sans alerte", "Certain (> 1/2)", "Aucun moyen de détection")]
    for i, rw in enumerate(rows, 4):
        for j, v in enumerate(rw, 1): b.cell(i, j, v)
    body(b, 4, 13, 4); b.freeze_panes = "A4"
    aide_b = b.cell(15, 1, "[Barèmes indicatifs : à adapter au domaine (produit, processus, logiciel).]")
    aide_b.font = Font(name=FONT, italic=True, color="808080")

    se = wb.create_sheet("Seuils"); title(se, "Seuils de criticité (modifiables)")
    header(se, 2, ["Seuil", "Valeur", "Signification"], [30, 10, 50])
    for i, (k, v, m) in enumerate([("Seuil « à surveiller »", 100, "C ≥ valeur : orange"),
                                   ("Seuil « action obligatoire »", 200, "C ≥ valeur : rouge")], 3):
        se.cell(i, 1, k); se.cell(i, 2, v); se.cell(i, 3, m)
    body(se, 3, 4, 3)

    s = wb.create_sheet("Synthèse"); title(s, "Synthèse AMDEC")
    A = "AMDEC!"
    J, Q = f"{A}J4:J{last}", f"{A}Q4:Q{last}"
    header(s, 3, ["Indicateur", "Avant action", "Après action"], [42, 15, 15])
    lines = [("Modes de défaillance cotés", f"=COUNT({J})", f"=COUNT({Q})"),
             ("Criticité maximale", f"=IF(COUNT({J})=0,0,MAX({J}))", f"=IF(COUNT({Q})=0,0,MAX({Q}))"),
             ("Criticité moyenne", f"=IF(COUNT({J})=0,0,AVERAGE({J}))", f"=IF(COUNT({Q})=0,0,AVERAGE({Q}))"),
             ("Nb critiques (≥ seuil action)", f"=COUNTIF({J},\">=\"&Seuils!B4)", f"=COUNTIF({Q},\">=\"&Seuils!B4)"),
             ("Nb à surveiller", f"=COUNTIFS({J},\">=\"&Seuils!B3,{J},\"<\"&Seuils!B4)", f"=COUNTIFS({Q},\">=\"&Seuils!B3,{Q},\"<\"&Seuils!B4)"),
             ("Nb acceptables", f"=COUNTIF({J},\"<\"&Seuils!B3)", f"=COUNTIF({Q},\"<\"&Seuils!B3)"),
             ("Somme des criticités", f"=SUM({J})", f"=SUM({Q})")]
    for i, rw in enumerate(lines, 4):
        for j, v in enumerate(rw, 1): s.cell(i, j, v)
    s["B6"].number_format = s["C6"].number_format = "0.0"
    s.cell(11, 1, "Réduction globale de criticité")
    s.cell(11, 2, "=IF(B10=0,0,(B10-C10)/B10)").number_format = "0%"
    s.cell(12, 1, "Critiques sans action définie")
    s.cell(12, 2, f'=COUNTIFS({J},">="&Seuils!B4,{A}K4:K{last},"")')
    body(s, 4, 12, 3); s.freeze_panes = "A4"
    mode_emploi(wb, ["Objet :", "Analyse des Modes de Défaillance, de leurs Effets et de leur Criticité.",
        "Saisie (onglet AMDEC) :", "1. Une ligne par couple mode de défaillance / cause.",
        "2. Coter G, O, D de 1 à 10 selon l'onglet Barèmes (saisie contrôlée).",
        "3. Criticité = G × O × D (calculée, colorée selon l'onglet Seuils).",
        "4. Pour toute criticité orange/rouge, définir une action, un responsable, une échéance.",
        "5. Après mise en œuvre, recoter G', O', D' : la criticité résiduelle et le gain se calculent.",
        "Lecture :", "Onglet Synthèse : répartition avant/après, réduction globale, critiques sans action.",
        "Rappel :", "La gravité G ne baisse en général que par une modification de conception."])
    wb.save(OUT / "amdec.xlsx")

if __name__ == "__main__":
    build_cdc(); build_revue(); build_trace(); build_amdec()
    print("Fichiers générés dans", OUT)
